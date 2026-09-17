#!/usr/bin/env python3
"""
log_tracker_crosscheck.py - reconcile one day's signals log against the MASTER.

Claude Equity Bot - Phase 3-C - advisory-only - no orders - not financial advice.

WHAT THIS IS
    The daily 24/24 cross-check (Daily_Workflow_v5_3 section 8), as a tool instead
    of a hand comparison. Compares every log row against the pasted row in
    logs\\claude_equity_bot_tracker.xlsx on: signal, confidence, data quality and
    price at signal.

READ MASTER / WRITE NOTHING
    Hard-guarded. Refuses to run against tracker_with_registry.xlsx, which is a
    GENERATED REPORT COPY and never an input. This tool writes nothing at all.

FORMULA CACHE
    Immune. Columns A-G (Date, Ticker, Account, Signal, Conf %, Data Quality,
    Price at Signal) are pasted raw values, not formulas. The Return/Right columns
    are never read here.

PRECISION
    Distinguishes a PRECISION finding (the log value rounded to the tracker's
    decimal places) from a MISMATCH (a genuinely different number). The GLD paste
    truncation was the former and would otherwise read as the latter.

USAGE
    python log_tracker_crosscheck.py --log logs\\signals_2026-07-30.log ^
        --tracker logs\\claude_equity_bot_tracker.xlsx
    python log_tracker_crosscheck.py --log ... --tracker ... --quiet   (exit 1 on findings)
"""

import argparse
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

FORBIDDEN_INPUTS = {"tracker_with_registry.xlsx"}
HEADER_ROW = 4          # data begins at row 5
COLS = {"date": 0, "ticker": 1, "account": 2, "signal": 3, "conf": 4, "dq": 5, "px": 6}

LINE_RE = re.compile(
    r"Signal JSON: (?P<json>\{.*?\})\s*\|\s*\[(?P<phase>PHASE[^\]]*)\]"
    r"\s*\[[^\]]*\]\s*\[[^\]]*\]\s*(?P<ticker>[A-Z.\-]+)\s*\|"
)
DATE_RE = re.compile(r"signals_(\d{4}-\d{2}-\d{2})", re.I)
TS_RE = re.compile(r"^(\d{4}-\d{2}-\d{2}) ")


def parse_log(path):
    """AM-canonical: first occurrence per ticker wins (Workflow section 9)."""
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    rows, seen, shadowed, phases = {}, set(), [], set()
    log_date = None
    for line in text.splitlines():
        if log_date is None:
            ts = TS_RE.match(line)
            if ts:
                log_date = datetime.strptime(ts.group(1), "%Y-%m-%d").date()
        m = LINE_RE.search(line)
        if not m:
            continue
        phases.add(m.group("phase"))
        t = m.group("ticker")
        if t in seen:
            shadowed.append(t)
            continue
        seen.add(t)
        try:
            p = json.loads(m.group("json"))
        except json.JSONDecodeError:
            continue
        rows[t] = {"signal": p.get("signal"), "conf": p.get("confidence"),
                   "dq": p.get("data_quality"), "px": p.get("lastPrice")}
    fn = DATE_RE.search(Path(path).name)
    if fn:
        fn_date = datetime.strptime(fn.group(1), "%Y-%m-%d").date()
        if log_date and fn_date != log_date:
            print("WARN: filename date {} != first timestamp date {}".format(fn_date, log_date))
        log_date = log_date or fn_date
    return rows, log_date, shadowed, sorted(phases)


def load_master(path, target):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Signals" not in wb.sheetnames:
        raise SystemExit("ERROR: no 'Signals' sheet in {}".format(path))
    ws = wb["Signals"]
    rows, dupes = {}, []
    for excel_row, r in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True),
                                  start=HEADER_ROW + 1):
        d = r[COLS["date"]]
        if d is None:
            continue
        d = d.date() if isinstance(d, datetime) else d
        if d != target:
            continue
        t = r[COLS["ticker"]]
        if t in rows:
            dupes.append((t, excel_row))
            continue
        rows[t] = {"row": excel_row, "signal": r[COLS["signal"]], "conf": r[COLS["conf"]],
                   "dq": r[COLS["dq"]], "px": r[COLS["px"]], "account": r[COLS["account"]]}
    wb.close()
    return rows, dupes


def previous_cohort_size(path, target):
    """Row count of the most recent cohort BEFORE `target`, or None.

    ADDED 2026-09-08 to replace the hardcoded --expect 24. The live universe is
    whatever the last session actually ran; comparing today against that answers
    "did the cohort change?", which is the question, and cannot go stale in a
    source file the way a copied constant does.
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Signals"]
        counts = {}
        for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            d = r[COLS["date"]]
            if d is None:
                continue
            d = d.date() if isinstance(d, datetime) else d
            if d < target:
                counts[d] = counts.get(d, 0) + 1
        wb.close()
        if not counts:
            return None, None
        prev = max(counts)
        return counts[prev], prev
    except Exception:
        return None, None


def decimals(value):
    try:
        exp = Decimal(str(value)).as_tuple().exponent
        return -exp if isinstance(exp, int) and exp < 0 else 0
    except Exception:
        return 0


def compare(log_rows, master_rows, conf_tol):
    findings = []
    for t in sorted(set(log_rows) | set(master_rows)):
        lg, ms = log_rows.get(t), master_rows.get(t)
        if lg and not ms:
            findings.append(("MISSING", t, "--", "in log, absent from master"))
            continue
        if ms and not lg:
            findings.append(("EXTRA", t, ms["row"], "in master, absent from log"))
            continue

        if lg["signal"] != ms["signal"]:
            findings.append(("MISMATCH", t, ms["row"],
                             "signal log {} vs master {}".format(lg["signal"], ms["signal"])))

        # Signals conf column stores fractions (0.52); normalise, never hand-convert.
        if ms["conf"] is None or lg["conf"] is None:
            findings.append(("MISMATCH", t, ms["row"], "confidence missing on one side"))
        else:
            master_pct = float(ms["conf"]) * 100.0
            if abs(master_pct - float(lg["conf"])) > conf_tol:
                findings.append(("MISMATCH", t, ms["row"],
                                 "conf log {}% vs master {:.4g}%".format(lg["conf"], master_pct)))

        if lg["dq"] != ms["dq"]:
            findings.append(("MISMATCH", t, ms["row"],
                             "DQ log {} vs master {}".format(lg["dq"], ms["dq"])))

        if ms["px"] is None or lg["px"] is None:
            findings.append(("MISMATCH", t, ms["row"], "price missing on one side"))
        elif abs(float(ms["px"]) - float(lg["px"])) > 1e-9:
            dp = decimals(ms["px"])
            if abs(round(float(lg["px"]), dp) - float(ms["px"])) < 1e-9:
                findings.append(("PRECISION", t, ms["row"],
                                 "px log {} pasted as {} ({} dp) - verify against the log line"
                                 .format(lg["px"], ms["px"], dp)))
            else:
                findings.append(("MISMATCH", t, ms["row"],
                                 "px log {} vs master {}".format(lg["px"], ms["px"])))
    return findings


def main():
    ap = argparse.ArgumentParser(description="Reconcile a signals log against the MASTER tracker.")
    ap.add_argument("--log", required=True, help=r"logs\signals_YYYY-MM-DD.log")
    ap.add_argument("--tracker", default=r"logs\claude_equity_bot_tracker.xlsx",
                    help=r"THE MASTER: logs\claude_equity_bot_tracker.xlsx")
    ap.add_argument("--date", help="override the cohort date (YYYY-MM-DD)")
    ap.add_argument("--conf-tol", type=float, default=0.05,
                    help="confidence tolerance in percentage points (default 0.05)")
    # ANALYST NOTE (2026-09-08): this default was a hardcoded 24 -- a hand-copied
    #   snapshot of a watchlist that moves. The live cohort has been 28 since at
    #   least 2026-09-04, so every run printed "rows: log 28 / master 28
    #   (expected 24)" plus a [COUNT] line reporting a mismatch that was not one.
    #   Harmless at 0 findings, but it is a SECOND SOURCE OF TRUTH for the
    #   universe -- the exact defect DP-3C-d unfroze at the EDGAR boundary,
    #   where the same 24-name constant survives only as a documented fallback
    #   with "CONFIRM it still matches the live watchlist" attached to it.
    #
    #   None now means DERIVE: the expectation becomes the PREVIOUS cohort's
    #   row count in the master, so the check answers the question actually
    #   worth asking -- did the universe change since the last session? -- and
    #   cannot go stale in a source file. An explicit --expect still pins it.
    ap.add_argument("--expect", type=int, default=None,
                    help="expected row count. Default: DERIVED from the "
                         "previous cohort in the master, so the check tracks "
                         "the live universe rather than a frozen constant.")
    ap.add_argument("--quiet", action="store_true", help="suppress stdout; exit 1 on findings")
    args = ap.parse_args()

    tracker = Path(args.tracker)
    if tracker.name.lower() in FORBIDDEN_INPUTS:
        print("REFUSED: {} is a GENERATED REPORT COPY and is never an input.\n"
              "         Point this at logs\\claude_equity_bot_tracker.xlsx (THE MASTER)."
              .format(tracker.name), file=sys.stderr)
        return 2
    if not tracker.is_file():
        print("ERROR: no such workbook: {}".format(tracker), file=sys.stderr)
        return 2
    if not Path(args.log).is_file():
        print("ERROR: no such log file: {}".format(args.log), file=sys.stderr)
        return 2

    log_rows, log_date, shadowed, phases = parse_log(args.log)
    target = datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else log_date
    if target is None:
        print("ERROR: could not determine the cohort date; pass --date", file=sys.stderr)
        return 2

    master_rows, dupes = load_master(tracker, target)
    findings = compare(log_rows, master_rows, args.conf_tol)

    expect, expect_src = args.expect, "given on the command line"
    if expect is None:
        expect, prev_date = previous_cohort_size(tracker, target)
        expect_src = (f"derived from the previous cohort {prev_date}"
                      if expect is not None else None)
    if expect is None:
        expect_src = "no prior cohort in the master -- count not checked"

    out = []
    out.append("=" * 78)
    out.append("LOG <-> MASTER CROSS-CHECK - {}".format(target))
    out.append("=" * 78)
    out.append("  log     : {}".format(Path(args.log).resolve()))
    out.append("  master  : {}".format(tracker.resolve()))
    out.append("  rows    : log {} / master {} (expected {} - {})".format(
        len(log_rows), len(master_rows),
        expect if expect is not None else "n/a", expect_src))
    out.append("  phase   : {}".format(", ".join(phases) if phases else
                                       "NONE MATCHED - no signal line parsed"))
    if shadowed:
        out.append("  [WARN] PM append ignored for: {}".format(", ".join(sorted(set(shadowed)))))
    for t, row in dupes:
        out.append("  [FAIL] duplicate master row for {} at excel row {}".format(t, row))
    out.append("")
    # FIXED 2026-09-16 (DP-3C-ae (ii), finishing the 2026-09-08 change). The
    # clean test and the exit code still compared against args.expect, which
    # is None whenever the expectation is DERIVED -- so every default run
    # printed "RESULT: 0 finding(s)" and exited 1, and the exit code carried no
    # signal again. Now: a log/master row-count mismatch for the same date is
    # a FAIL (the reconciliation this tool exists for); a change against the
    # previous cohort is a WARN (watchlist churn) and never sets the exit code.
    count_fail = len(log_rows) != len(master_rows)
    if not findings and not dupes and not count_fail:
        out.append("  RESULT: {}/{} CLEAN - signal, confidence, DQ and price all reconcile."
                   .format(len(log_rows), len(master_rows)))
        if expect is not None and len(log_rows) != expect:
            out.append("  [WARN     ] log has {} rows, {} had {} - the universe "
                       "CHANGED (informational; does not set the exit code)"
                       .format(len(log_rows), expect_src, expect))
    else:
        out.append("  RESULT: {} finding(s)".format(len(findings) + len(dupes)))
        for kind, t, row, detail in findings:
            out.append("  [{:<9}] {:<6} excel row {:<6} {}".format(kind, t, row, detail))
        if count_fail:
            out.append("  [COUNT    ] log has {} rows, master has {} for the same date"
                       .format(len(log_rows), len(master_rows)))
        if expect is not None and len(log_rows) != expect:
            out.append("  [WARN     ] log has {} rows, {} had {} - the universe "
                       "CHANGED (429 skips log and are inert; informational)"
                       .format(len(log_rows), expect_src, expect))
    out.append("")
    out.append("Advisory-only - no orders - not financial advice.")

    if not args.quiet:
        print("\n".join(out))
    return 1 if (findings or dupes or count_fail) else 0


if __name__ == "__main__":
    sys.exit(main())
