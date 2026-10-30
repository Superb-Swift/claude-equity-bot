#!/usr/bin/env python3
"""
value_card.py -- the valuation layer, organised around decisions and changes.

WHY (2026-10-14, claude/ValuePass_WhatItMeasures_2026-10-14.md)
    implied_return prints about 27 KB a day, and across the 29 dated exhibits
    from 2026-08-24 to 10-14, 14 of 15 headline classes never changed. Three
    findings explain why, and each section below answers one of them:

    1. THE CLASS IS A TRAILING-FISCAL-YEAR P/E. Book value cancels out of
       `implied/trailing @9%`:
           ratio = (P/EPS) * (r-g)/(1+g) + g/((1+g) ROE)  =  0.0583 * P/E + 0.0291/ROE
       Schwab's eps is the last fiscal year and does not move on a quarterly
       report, so the class moves only with price or a 10-K. Section 6 checks
       this identity every run and FLAGS it if the exhibit stops satisfying it.
    2. THE DAILY REPRINT HIDES THE DAYS THAT MATTER. Section 1 prints only
       what changed since the previous dated exhibit, and splits every ratio
       move into a PRICE part and a FILING part.
    3. THE CLASSES ARE FRAGILE. Neither ROE route is trailing-twelve-month.
       Section 3 re-derives every class under both routes and r = 7/9/11%,
       and calls a class ROBUST only if all six agree.

    Section 2 is the decision card the 2026-08-24 DP-V2 memo specified as
    stage 1 and the 2026-10-09 review ranked V1. Section 4 is DP-V2's
    counterfactual flip ledger: COUNTS ONLY, never an outcome.

DP-V4 (2026-10-24, claude/DP-V4_ValueCard_TTM_2026-10-24.md): THE CLASS
    READS 12-MONTH ROE. Finding 1 above was the defect: Schwab's last-FY eps
    is 1 to 4 quarters old (NVDA's 4.90 against 7.91 over the last 12 months),
    so the class could not hear a quarterly report. It is now
        implied ROE at 9% / (12-month net income / period-end equity)
    from sec_facts, POINT IN TIME: a class dated d reads only filings SEC had
    published by d (on or before d for the card; strictly before d for the
    DP-V2 ledger, whose BUY was decided that morning). The 2.00 / 1.50 lines,
    r = 9% and g = 3% are unchanged. Robustness is the r = 7/9/11% grid on that
    basis plus the distance to the nearest line; the last-FY class is printed
    beside it for reference. The DP-V2 ledger is REBUILT on the 12-month basis
    from 2026-08-24; the last-FY ledger is CLOSED and printed, never pooled.
    A name with no 12-month figure is NO READ, never a last-FY fallback.
    The card also opens the hood, one question per panel:
      2b  earnings direction, earnings quality, balance sheet (SEC facts)
      2c  expectations: the next report, Finnhub's ADJUSTED consensus and the
          last four surprises (earnings_calendar), and the class if the
          consensus is met (approximate; its own basis, never the class)

2b, 2026-10-29 (claude/UnderTheHood_GapAnalysis_2026-10-29.md, step 1): the
    cells the 10-28 sheet left blank for a reason it did not print. OCF/rev
    (operating cash flow / revenue) is defined through a loss, where OCF/NI is
    not; L/A (liabilities / assets) is defined when equity is negative, where
    L/E is not; the balance sheet opens on the latest assets date when there
    is no equity line (PAA's cash). The text card and the sheet share one note
    list (hood_notes) and one +/-300% cap. Display only: the class, its
    robustness and the DP-V2 ledger are unchanged.

WHAT IT DOES NOT DO
    It reads no return, breach, hit or +Nd cell. It modifies no gate, admits
    and rejects nothing, touches no workbook, no prompt, no endpoint, no
    edgar_client extract. DP-V1 stands: nothing here reaches the model. The
    DP-V2 gate column is COUNTERFACTUAL and is never applied.

INPUTS
    implied_return_<date>.md   today's, plus every earlier dated exhibit found
                               in the root or exhibits\\archive (sections 1, 4)
    logs\\signals_<date>.log   signal, confidence, held/scout, per-account
    quote_dump.json            divYield and eps for the book-free reads (5)
    the MASTER                 BUY rows since 2026-08-24 for the ledger (4).
                               tracker_with_registry.xlsx is REFUSED.

USAGE
    python value_card.py --date 2026-10-14 --tracker logs\\claude_equity_bot_tracker.xlsx
           --log logs\\signals_2026-10-14.log --quotes quote_dump.json
           --md value_card_2026-10-14.md [--json value_card_2026-10-14.json]
    python value_card.py --selftest

EXIT  0 = exhibit written.  1 = written, with FINDINGS (identity check failed,
      or a parse anomaly).  2 = could not run (no exhibit for the date, report
      copy given as input, unreadable master or log). Nothing is written on 2.

Advisory-only - no orders - not financial advice.
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import json
import os
import re
import sys

R_GRID = (7.0, 9.0, 11.0)      # cost-of-equity grid printed by implied_return 4
R_BASE = 9.0                   # the pre-registered class rate (DP-V2 memo 3.1)
G_BASE = 3.0                   # implied_return's g
DEMANDING_AT = 2.00            # DP-V2 memo 3.1, frozen at first run, never re-cut
NEUTRAL_AT = 1.50
DPV2_GATE = {"DEMANDING": 75, "NEUTRAL": 70, "UNDEMANDING": 65, "NO READ": 70}
DPV2_BASE_GATE = 70
DPV2_SINCE = dt.date(2026, 8, 24)      # first day the layer printed section 4
DPV2_STOP = dt.date(2026, 12, 31)      # DP-V2 memo 6: withdrawn, not extended
DPV2_MIN_ROWS, DPV2_MIN_TICKERS = 12, 3
INSTALLED = dt.date(2026, 10, 14)      # rows on or before this are BACKFILLED
DPV4_ON = dt.date(2026, 10, 24)        # DP-V4: the class reads 12-month ROE; earlier rows REBUILT
EDGE_BORDERLINE = 0.05                 # a ratio this close to a line can cross on a 2-3% price move
CLASS_LETTER = {"DEMANDING": "D", "NEUTRAL": "N", "UNDEMANDING": "U", "NO READ": "-"}
# The identity is checked against the ROUNDING of the printed inputs (px, BVPS,
# ROE% and the ratio are all printed to 2dp), so SND's 0.51% ROE gets the wide
# bound it needs and a 2.7x name gets a tight one. Verified on all 29 exhibits
# 2026-08-24..10-14: no name exceeds its bound.
REFUSED = "tracker_with_registry.xlsx"
SECTION_RULE = re.compile(r"^-{20,}")

ACCT_RE = re.compile(r"\] \[([^\]]+)\] \[(HELD|SCOUT)\] ([A-Z][A-Z0-9.]*) \| Signal: (\w+) \| "
                     r"Confidence: (\d+)% \| Risk: (\w+)")
PER_RE = re.compile(r"PER-ACCOUNT: ([^.|]+)")


# --------------------------------------------------------------------------- #
# Parsing the implied_return exhibit (format stable since 2026-08-24)
# --------------------------------------------------------------------------- #

def _num(s):
    s = str(s).strip().rstrip("%x").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _section(lines, title_re):
    """The body lines of the section whose title matches title_re."""
    out, on = [], False
    for i, line in enumerate(lines):
        if SECTION_RULE.match(line) and i + 1 < len(lines) and re.match(title_re, lines[i + 1]):
            on = True
            continue
        if on and SECTION_RULE.match(line) and out and not re.match(title_re, lines[i - 1]):
            break
        if on:
            out.append(line)
    return out


S1_RE = re.compile(
    r"^\s+([A-Z][A-Z0-9.]*)\s+(\d{4}-\d{2}-\d{2})\s+(-?[\d.]+)\s+(-?[\d.]+)\s+([+-][\d.]+%)\s+"
    r"(n/a|-?[\d.]+%)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(?:([\d.]+)Q\s+([\d.]+)Q\s+)?(.*)$")
S1A_RE = re.compile(r"^\s+([A-Z][A-Z0-9.]*)\s+(Q\d|FY)\s+(n/a|-?[\d.]+%)\s+(n/a|-?[\d.]+%)\s+")
S4_RE = re.compile(r"^\s+([A-Z][A-Z0-9.]*)\s+(-?[\d.]+)\s+(-?[\d.]+%|n/a)\s+(-?[\d.]+)%\s+"
                   r"(-?[\d.]+)%\s+(-?[\d.]+)%\s+(n/a|-?[\d.]+x)")
S4B_RE = re.compile(r"^\s+([A-Z][A-Z0-9.]*)\s+(-?[\d.]+)\s+(-?[\d.]+)\s+(n/a|-?[\d.]+%)\s+"
                    r"(NO CONVERGENT g|-?[\d.]+%)\s+(.*)$")
BLOCK_RE = re.compile(r"^\s+([A-Z][A-Z0-9.]*) NOT COMPUTED -- (.*)$")
DATE_RE = re.compile(r"implied_return_(\d{4}-\d{2}-\d{2})\.md$")


def parse_exhibit(path):
    """{ticker: fields} and {ticker: blocked reason} from one implied_return exhibit."""
    with open(path, encoding="utf-8", errors="replace") as fh:
        lines = fh.read().replace("\r", "").split("\n")
    rec, blocked = {}, {}
    for line in _section(lines, r"^1\. INPUTS"):
        m = S1_RE.match(line)
        if m:
            rec[m.group(1)] = dict(
                anchor=m.group(2), bvps=_num(m.group(3)), px=_num(m.group(4)),
                roe=None if m.group(6) == "n/a" else _num(m.group(6)),
                mb=_num(m.group(7)), divy=_num(m.group(8)),
                anc_age=_num(m.group(9)) if m.group(9) else None,
                num_age=_num(m.group(10)) if m.group(10) else None,
                basis=m.group(11).strip())
    for line in _section(lines, r"^1a\."):
        m = S1A_RE.match(line)
        if m and m.group(1) in rec:
            rec[m.group(1)]["roe_edgar"] = None if m.group(4) == "n/a" else _num(m.group(4))
    for line in _section(lines, r"^4\. \(II\)"):
        m = S4_RE.match(line)
        if m and m.group(1) in rec:
            rec[m.group(1)].update(
                i7=_num(m.group(4)), i9=_num(m.group(5)), i11=_num(m.group(6)),
                ratio=None if m.group(7) == "n/a" else _num(m.group(7)))
    for line in _section(lines, r"^4b\."):
        m = S4B_RE.match(line)
        if m and m.group(1) in rec:
            rec[m.group(1)]["g_impl"] = None if "NO" in m.group(5) else _num(m.group(5))
    for line in lines:
        m = BLOCK_RE.match(line)
        if m:
            blocked[m.group(1)] = m.group(2).strip()
    return rec, blocked


def find_exhibits(root):
    """{date: path} for every dated implied_return exhibit in root and exhibits/archive."""
    out = {}
    for p in glob.glob(os.path.join(root, "implied_return_*.md")) + \
            glob.glob(os.path.join(root, "exhibits", "archive", "implied_return_*.md")):
        m = DATE_RE.search(os.path.basename(p))
        if m:
            d = dt.date.fromisoformat(m.group(1))
            if d not in out or os.path.dirname(p) == root:   # root copy wins
                out[d] = p
    return out


# --------------------------------------------------------------------------- #
# Valuation arithmetic -- the exhibit's own identity, nothing new
# --------------------------------------------------------------------------- #

def klass(ratio):
    if ratio is None:
        return "NO READ"
    if ratio >= DEMANDING_AT:
        return "DEMANDING"
    return "NEUTRAL" if ratio >= NEUTRAL_AT else "UNDEMANDING"


def implied_roe(mb, r=R_BASE, g=G_BASE):
    """implied_return (II), percent in and out: [r + (M/B - 1)(r - g)] / (1 + g),
    i.e. [M/B (r - g) + g] / (1 + g). Same expression as implied_return.py L179."""
    return (mb * (r - g) + g) / (1.0 + g / 100.0)


def ratio_from(px, bvps, roe, r=R_BASE):
    """implied/trailing at rate r. None where trailing ROE is not positive."""
    if not roe or roe <= 0 or not bvps or bvps <= 0 or not px:
        return None
    return implied_roe(px / bvps, r) / roe


def pe_lastfy(f):
    """Price over the EPS the class actually uses (ROE x BVPS)."""
    if not f.get("roe") or f["roe"] <= 0 or not f.get("bvps"):
        return None
    return f["px"] / (f["roe"] / 100.0 * f["bvps"])


def robustness(f):
    """DP-V4 (2026-10-24): the class under the r = 7/9/11% grid on the 12-month
    basis, and the distance from today's ratio to the nearest line. ROBUST
    only if the three agree. The old six-way test mixed in two ROE routes that
    were both known to be the wrong period (the last-FY year and one quarter
    x4); the basis is now settled by DP-V4, so what is left to doubt is the
    cost of equity and the closeness of a line. The last-FY class (the basis
    before DP-V4) is carried for reference only."""
    roe = f.get("roe")
    combos = {}
    for r, ik in zip(R_GRID, ("i7", "i9", "i11")):
        i = f.get(ik)
        ratio = (i / roe) if (i is not None and roe and roe > 0) else None
        combos[r] = (ratio, klass(ratio))
    base = klass(f.get("ratio"))
    ratios = [v[0] for v in combos.values() if v[0] is not None]
    classes = {v[1] for v in combos.values()}
    x = f.get("ratio")
    edge = min((abs(x - line), line) for line in (NEUTRAL_AT, DEMANDING_AT)) if x is not None else None
    lastfy = klass(f.get("ratio_lastfy")) if "ratio_lastfy" in f else None
    return dict(
        base=base,
        tag=("ROBUST" if len(classes) == 1 and base != "NO READ" else
             ("—" if base == "NO READ" else "FRAGILE")),
        lo=min(ratios) if ratios else None, hi=max(ratios) if ratios else None,
        r_flip=(base != "NO READ" and len(classes) > 1),
        r_classes=[combos[r][1] for r in R_GRID],
        edge=edge[0] if edge else None, edge_line=edge[1] if edge else None,
        borderline=bool(edge) and edge[0] < EDGE_BORDERLINE,
        lastfy=lastfy, lastfy_ratio=f.get("ratio_lastfy"),
        basis_flip=(lastfy is not None and lastfy != base))


def decompose(prev, cur):
    """Split a ratio move into PRICE (new price, old fundamentals) and FILING
    (new fundamentals at the new price). Returns None where either side has
    no ratio."""
    r0 = ratio_from(prev["px"], prev["bvps"], prev["roe"])
    rp = ratio_from(cur["px"], prev["bvps"], prev["roe"])
    r1 = ratio_from(cur["px"], cur["bvps"], cur["roe"])
    if None in (r0, rp, r1):
        return None
    return dict(r0=r0, r1=r1, price=rp - r0, filing=r1 - rp)


def g_implied_pct(roe_pct, mb, r=R_BASE):
    """implied_return (III) on the class's ROE, percent: the growth the price
    implies, g = (M/B r - ROE) / (M/B + ROE - 1). None when no convergent g
    exists (g >= r) or M/B is not positive -- implied_return.g_implied."""
    if roe_pct is None or mb is None or mb <= 0:
        return None
    roe, rs = roe_pct / 100.0, r / 100.0
    den = mb + roe - 1.0
    if den == 0:
        return None
    g = (mb * rs - roe) / den
    return None if g >= rs else g * 100.0


# --------------------------------------------------------------------------- #
# DP-V4: the 12-month basis, point in time, from the slim SEC cache
# --------------------------------------------------------------------------- #

class TTMBasis:
    """12-month ROE and the fundamentals panels per ticker and date, from
    valuation_cache\\ (written by ttm_eps / sec_data_audit; read only here).
    Each filer is loaded once; every read is point in time (sec_facts)."""

    def __init__(self, root):
        import sec_facts as SF
        self.SF, self.root = SF, root
        self.cik = SF.load_saved_cik_map(root)
        self._cf = {}

    def cf(self, t):
        if t not in self._cf:
            self._cf[t] = self.SF.load_slim(t, self.cik, self.root)
        return self._cf[t]

    def why_none(self, t):
        if not self.cik:
            return "no 12-month figure: valuation_cache\\cik_map.json missing (run ttm_eps)"
        if t not in self.cik:
            return "no 12-month figure: not an SEC filer in the cache map"
        return "no 12-month figure: no slim cache for this filer (run ttm_eps)"

    def roe(self, t, d, strict):
        """(roe percent or None, ttm_roe dict or None, why or None)."""
        cf = self.cf(t)
        if cf is None:
            return None, None, self.why_none(t)
        r = self.SF.ttm_roe(self.SF.point_in_time(cf, d, strict=strict), d)
        pct = r["roe"] * 100.0 if r["roe"] is not None else None
        return pct, r, r["why"]

    def fundamentals(self, t, d):
        cf = self.cf(t)
        if cf is None:
            return None
        return self.SF.fundamentals(self.SF.point_in_time(cf, d, strict=False), d)


def rebase(rec, basis, d, strict):
    """The exhibit's records re-based on 12-month ROE as of d (DP-V4).
    `roe` and `ratio` become the 12-month figures; the exhibit's own are kept
    as roe_lastfy / ratio_lastfy for reference and for the identity check.
    The implied ROE at each r (i7/i9/i11) depends only on price and book, so
    it carries over unchanged: ratio = i9 / 12-month ROE."""
    out = {}
    for t, f in rec.items():
        g = dict(f, roe_lastfy=f.get("roe"), ratio_lastfy=f.get("ratio"),
                 g_impl_lastfy=f.get("g_impl"))
        pct, r, why = basis.roe(t, d, strict) if basis is not None else (None, None, "no basis")
        g["roe"] = pct
        g["ttm_why"] = why
        g["ttm_notes"] = (r or {}).get("notes") or []
        g["ttm_end"] = (r or {}).get("ni_end")
        i9 = f.get("i9")
        g["ratio"] = (i9 / pct) if (i9 is not None and pct is not None and pct > 0) else None
        g["g_impl"] = g_implied_pct(pct, f.get("mb")) if pct is not None else None
        out[t] = g
    return out


def no_read_reason(f, blocked_reason=None):
    if blocked_reason:
        return blocked_reason
    if f is None:
        return "not in the valuation set"
    return f.get("ttm_why") or "negative 12-month ROE"


def sector_caveat(t):
    """A one-line warning where a panel's ratio does not mean what it means
    elsewhere. Read from sector_map; nothing is computed from it."""
    try:
        import sector_map
        s = sector_map.STOCK_SECTOR.get(t)
    except Exception:                                   # noqa: BLE001
        s = None
    if s == "XLF":
        return "bank: operating cash flow and liabilities/equity are not comparable"
    if s == "XLRE":
        return "REIT: GAAP earnings carry depreciation; funds from operations are not read"
    return ""


def pct_s(x, cap=3.0):
    """'+12%'; beyond +/-300% the base is too small to read as a rate."""
    if x is None:
        return "—"
    if abs(x) > cap:
        return ">+300%" if x > 0 else "<-300%"
    return "%+.0f%%" % (100.0 * x)


def pct0(x):
    """'9%' for a ratio 0.09; '—' when undefined."""
    return "—" if x is None else "%.0f%%" % (100.0 * x)


def _p100(x):
    return None if x is None else 100.0 * x


def growth_cell(x, cap=3.0):
    """A sheet cell for a y/y change in percent; beyond +/-300% the same text
    the card prints (2026-10-29: the sheet printed SND 1200 and DOW 675)."""
    if x is None:
        return None
    if abs(x) > cap:
        return ">+300%" if x > 0 else "<-300%"
    return 100.0 * x


def hood_notes(u):
    """The 2b notes for one name, ONE list for the text card and the sheet
    (2026-10-29: the sheet had dropped 'negative book equity'). `u` is the
    fundamentals dict plus its sector caveat."""
    b, q = u.get("balance") or {}, u.get("quality") or {}
    out = [u["caveat"]] if u.get("caveat") else []
    if b.get("negative_equity"):
        eq = b.get("equity")
        out.append("negative book equity%s: L/E not defined, see L/A"
                   % ("" if eq is None else " (deficit $%.2fB)" % (abs(eq) / 1e9)))
    if q.get("net_loss"):
        out.append("12-month net loss: OCF/NI not defined, see OCF/rev")
    rn = roic_note_for(u)
    return out + list(u.get("notes") or []) + ([rn] if rn else [])


def roic_note_for(u):
    """DP-V6: the ROIC note for one name ('' for a standard read; a bank's ROIC
    is not shown, and its caveat already says why)."""
    if (u.get("caveat") or "").startswith("bank"):
        return ""
    try:
        import sec_facts as SF
        return SF.roic_note(u.get("roic"))
    except Exception:                                   # noqa: BLE001
        return ""


def roic_value(u):
    """DP-V6: the ROIC fraction to show, or None (no read, or a bank)."""
    if (u.get("caveat") or "").startswith("bank"):
        return None
    return (u.get("roic") or {}).get("roic")


def hood_sheet_row(t, u):
    """The 2b sheet row, in HOOD_SHEET_COLUMNS order."""
    g, q, b = u["growth"], u["quality"], u["balance"]
    bank = (u.get("caveat") or "").startswith("bank")
    return [t, g["eps"]["ttm"], growth_cell(g["eps"]["chg"]), growth_cell(g["revenue"]["chg"]),
            growth_cell(g["q_eps"]["chg"]), None if bank else q.get("ocf_to_ni"),
            None if bank else _p100(q.get("ocf_margin")),
            _p100(q.get("gross_margin")), _p100(q.get("gross_margin_year_ago")),
            _p100(q.get("op_margin")), _p100(q.get("op_margin_year_ago")),
            _p100(roic_value(u)),
            None if bank else b.get("liab_to_equity"),
            None if bank else _p100(b.get("liab_to_assets")),
            None if b.get("cash") is None else b["cash"] / 1e9,
            "; ".join(hood_notes(u))]


HOOD_SHEET_COLUMNS = (
    "Ticker", "EPS 12m", "EPS y/y %", "Revenue y/y %", "Last quarter EPS y/y %", "OCF / NI",
    "OCF / revenue %", "Gross margin %", "Gross margin a year ago %", "Operating margin %",
    "Operating margin a year ago %", "ROIC %", "Liabilities / equity", "Liabilities / assets %",
    "Cash ($B)", "Notes")


def dividend_implied_g(yield_pct, r=R_BASE):
    """Gordon, solved for g: P = D0(1+g)/(r-g)  ->  g = (r - y)/(1 + y)."""
    if yield_pct is None or yield_pct <= 0:
        return None
    y, rr = yield_pct / 100.0, r / 100.0
    return (rr - y) / (1.0 + y) * 100.0


# --------------------------------------------------------------------------- #
# Inputs other than the exhibit
# --------------------------------------------------------------------------- #

def parse_log(path):
    """{ticker: dict(acct, held, signal, conf, risk, per)} from a signals log."""
    out = {}
    with open(path, encoding="utf-8", errors="replace") as fh:
        for line in fh:
            m = ACCT_RE.search(line)
            if not m:
                continue
            per = PER_RE.search(line)
            per_map = {}
            if per:
                for part in per.group(1).split(";"):
                    if "=" in part:
                        k, v = part.split("=", 1)
                        per_map[k.strip()] = v.strip()
            px = None
            if "Signal JSON:" in line:
                try:
                    j, _ = json.JSONDecoder().raw_decode(line.split("Signal JSON:", 1)[1].strip())
                    px = j.get("lastPrice")
                except ValueError:
                    px = None
            out[m.group(3)] = dict(acct=m.group(1).strip(), held=m.group(2) == "HELD",
                                   signal=m.group(4).upper(), conf=int(m.group(5)),
                                   risk=m.group(6).upper(), per=per_map, px=px)
    return out


def executability(s):
    if s["signal"] == "HOLD":
        return "—"
    if not s["per"]:
        return "BLOCKED" if s["risk"] != "APPROVED" else "APPROVED"
    ok = [k for k, v in s["per"].items() if v == "APPROVED"]
    if len(ok) == len(s["per"]):
        return "EXECUTABLE"
    if ok:
        return "PARTIAL (" + ", ".join(ok) + " only)"
    return "REJECTED in every account"


def load_quotes(path):
    if not path or not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as fh:
        q = json.load(fh)
    return q.get("quotes", {}) if isinstance(q, dict) else {}


def load_buy_rows(tracker, since):
    """(date, ticker, conf) for BUY rows on/after `since`. Reads date, ticker,
    signal and confidence ONLY -- no price, return or outcome column is read."""
    import openpyxl
    from build_actionable_registry import (SIGNALS_SHEET, FIRST_DATA_ROW,
                                           LAST_DATA_ROW_CAP, COL, as_date, conf_pct)
    wb = openpyxl.load_workbook(tracker, data_only=True, read_only=True)
    ws = wb[SIGNALS_SHEET]
    out, blank = [], 0
    last = max(COL["date"], COL["ticker"], COL["signal"], COL["conf"])
    for vals in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=LAST_DATA_ROW_CAP,
                             min_col=1, max_col=last, values_only=True):
        d = as_date(vals[COL["date"] - 1])
        if d is None:
            blank += 1
            if blank >= 25:
                break
            continue
        blank = 0
        if d < since or str(vals[COL["signal"] - 1] or "").strip().upper() != "BUY":
            continue
        c = conf_pct(vals[COL["conf"] - 1])
        if c is not None:
            out.append((d, str(vals[COL["ticker"] - 1] or "").strip().upper(), c))
    wb.close()
    return out


# --------------------------------------------------------------------------- #
# DP-V2 counterfactual flip ledger -- counts only
# --------------------------------------------------------------------------- #

def dpv2_flip(cls, conf):
    """'ADMIT->REJECT', 'REJECT->ADMIT' or None, for a BUY at `conf` under DP-V2."""
    gate = DPV2_GATE.get(cls, DPV2_BASE_GATE)
    now, then = conf >= DPV2_BASE_GATE, conf >= gate
    if now and not then:
        return "ADMIT->REJECT"
    if then and not now:
        return "REJECT->ADMIT"
    return None


def _label_lastfy(d):
    return "backfilled" if d <= INSTALLED else "live"


def _label_ttm(d):
    """DP-V4: rows before the ruling are REBUILT point in time (only filings
    public before that session, the exhibit written that day); they count."""
    return "rebuilt" if d < DPV4_ON else "live"


def dpv2_ledger(buy_rows, classes_by_date, label=_label_lastfy):
    """classes_by_date: {date: {ticker: class}}. Returns flips, unscored rows.
    A row on a session with no exhibit is NEVER counted. It is reported as
    UNSCORED, with the class from the nearest EARLIER exhibit as a hint only,
    and only when that hint would have flipped it."""
    flips, unscored = [], []
    dates = sorted(classes_by_date)
    for d, tk, c in sorted(buy_rows):
        if d not in classes_by_date:
            prior = [x for x in dates if x < d]
            hint = classes_by_date[prior[-1]].get(tk, "NO READ") if prior else None
            if hint and dpv2_flip(hint, c):
                unscored.append((d, tk, c, hint, prior[-1]))
            continue
        cls = classes_by_date[d].get(tk, "NO READ")
        f = dpv2_flip(cls, c)
        if f:
            flips.append((d, tk, c, cls, f, label(d)))
    return flips, unscored


# --------------------------------------------------------------------------- #
# The exhibit
# --------------------------------------------------------------------------- #

def fmt(v, spec="{:.2f}", none="—"):
    return none if v is None else spec.format(v)


def load_context(root, date):
    """DP-V3 (2026-10-22): the TTM exhibit (ttm_eps.py, V2) and the earnings
    calendar (earnings_calendar.py, V5) for the SAME date, if value.bat wrote
    them before this step. Missing files mean '—', never a guess.
    AMENDED 2026-10-24 (DP-V4): the CLASS now reads 12-month ROE itself, from
    valuation_cache\\ (TTMBasis); these two files remain context: the notes,
    the report date, and the calendar's adjusted consensus (section 2c)."""
    out = {}
    for key, name in (("ttm", "ttm_eps_%s.json"), ("earn", "earnings_calendar_%s.json")):
        path = os.path.join(root, name % date)
        try:
            with open(path, encoding="utf-8") as fh:
                data = json.load(fh)
        except (OSError, ValueError):
            data = None
        if key == "earn" and isinstance(data, list):
            data = {r.get("ticker"): r for r in data}
        out[key] = data
    return out


THESIS_LABEL = ("THESIS BROKEN (Rob, 2026-10-22): off the scout list at the exit; "
                "until then, measurement only (DP-3C-bk)")


def thesis_label(t):
    """DP-3C-bk (2026-10-23): a name Rob removed from his watchlist because its
    thesis is broken, kept in the run until the exit for the measurement. The
    label tells the reader a BUY or SELL on it is not a candidate. Read from
    rulings.py, so it disappears when the ruling closes. Reporting only."""
    try:
        import rulings
        r = rulings.RULINGS.get("DP-3C-bk") or {}
    except Exception:                                   # noqa: BLE001
        return ""
    return THESIS_LABEL if (r.get("open") and t in (r.get("measurement_only") or ())) else ""


def _ctx_cells(ctx, t):
    """(ttm_eps, ttm_roe_pct, next_report, sessions_away, note) for ticker t.

    AMENDED 2026-10-23: the note joins every reason instead of showing only the
    first: the calendar flag, the TTM flags (STALE, negative equity, an equity
    date before the TTM end), a NO READ reason (without the doubled "TTM: no
    TTM:" of 10-22) and the DP-3C-bk thesis label."""
    tt = ((ctx.get("ttm") or {}).get(t)) or {}
    ee = ((ctx.get("earn") or {}).get(t)) or {}
    roe = tt.get("roe_ttm")
    bits = [ee["note"]] if ee.get("note") else []
    nr = tt.get("no_read")
    if nr and not ("fund" in nr and "fund" in (ee.get("note") or "")):
        bits.append(nr if nr.startswith("no TTM") else "TTM: " + nr)
    bits += ["TTM " + f for f in (tt.get("flags") or [])]
    if thesis_label(t):
        bits.append(thesis_label(t))
    return (tt.get("eps_ttm"), (roe * 100.0) if isinstance(roe, (int, float)) else None,
            ee.get("date"), ee.get("sessions"), "; ".join(bits))


def build(date, root, log_path, quotes_path, tracker):
    exhibits = find_exhibits(root)
    if date not in exhibits:
        raise SystemExit(("no implied_return_%s.md in %s or exhibits/archive -- run "
                          "implied_return first; nothing to build on." % (date, root), 2))
    rec_raw, blocked = parse_exhibit(exhibits[date])
    if not rec_raw:
        raise SystemExit(("implied_return_%s.md parsed to zero names -- format changed?" % date, 2))
    basis = TTMBasis(root)
    rec = rebase(rec_raw, basis, date, strict=False)        # the card: everything filed by today
    earlier = sorted(d for d in exhibits if d < date)
    prev_date = earlier[-1] if earlier else None
    prev_raw, prev_blocked = parse_exhibit(exhibits[prev_date]) if prev_date else ({}, {})
    prev_rec = rebase(prev_raw, basis, prev_date, strict=False) if prev_date else {}
    sig = parse_log(log_path)
    quotes = load_quotes(quotes_path)
    ctx = load_context(root, date)
    findings = []

    rob = {t: robustness(f) for t, f in rec.items()}
    held = sorted(t for t, s in sig.items() if s["held"])
    directional = sorted((t for t, s in sig.items() if s["signal"] in ("BUY", "SELL")),
                         key=lambda t: -sig[t]["conf"])
    if not basis.cik:
        findings.append("no valuation_cache\\cik_map.json: every class is NO READ until ttm_eps "
                        "or sec_data_audit has run online once (DP-V4)")

    L = []
    L.append("=" * 100)
    L.append("VALUE CARD -- the valuation layer organised around decisions and changes")
    L.append("Counterfactual only. Modifies no gate, admits and rejects nothing, reads no outcome.")
    L.append("=" * 100)
    L.append("as of %s   exhibit %s   previous exhibit %s   log %s" % (
        date, os.path.basename(exhibits[date]), prev_date or "none", os.path.basename(log_path)))
    L.append("class = implied ROE at 9%% / 12-MONTH ROE (DP-V4: SEC filings, point in time); "
             "lines %.2fx / %.2fx (DP-V2 memo 3.1, frozen)" % (DEMANDING_AT, NEUTRAL_AT))
    L.append("")

    # ---- 0. coverage -------------------------------------------------------
    read = [t for t in rec if rob[t]["base"] != "NO READ"]
    held_read = [t for t in held if t in rec and rob[t]["base"] != "NO READ"]
    held_rob = [t for t in held_read if rob[t]["tag"] == "ROBUST"]
    dir_read = [t for t in directional if t in rec and rob[t]["base"] != "NO READ"]
    L.append("-" * 100)
    L.append("0. COVERAGE -- stated first, because it bounds everything below")
    L.append("-" * 100)
    L.append("  watchlist %d   classed %d   ROBUST %d" % (
        len(sig), len(read), sum(1 for t in read if rob[t]["tag"] == "ROBUST")))
    L.append("  held names %d   classed %d   ROBUST %d   %s" % (
        len(held), len(held_read), len(held_rob),
        ("(" + ", ".join(held_rob) + ")") if held_rob else ""))
    L.append("  today's directionals %d   classed %d   %s" % (
        len(directional), len(dir_read), ", ".join(dir_read) or ""))
    coverage = dict(watchlist=len(sig), classed=len(read),
                    robust=sum(1 for t in read if rob[t]["tag"] == "ROBUST"),
                    held=len(held), held_classed=len(held_read), held_robust=held_rob,
                    directional=len(directional), directional_classed=dir_read)
    L.append("")

    # ---- 1. what changed ---------------------------------------------------
    L.append("-" * 100)
    L.append("1. WHAT CHANGED since %s -- filings, class crossings, robustness, coverage"
             % (prev_date or "(no earlier exhibit)"))
    L.append("-" * 100)
    basis_diff = [t for t in read if rob[t]["basis_flip"]] + \
           [t for t in rec if rob[t]["base"] == "NO READ" and rob[t]["lastfy"] not in (None, "NO READ")]
    basis_line = ("BASIS: 12-month ROE since %s (DP-V4). Today %d class(es) differ from the last-FY "
                  "basis: %s" % (DPV4_ON, len(basis_diff), ", ".join(
                      "%s %s->%s" % (t, rob[t]["lastfy"], rob[t]["base"]) for t in sorted(basis_diff))
                      or "none"))
    L.append("  " + basis_line)
    changes = []
    if prev_date:
        for t in sorted(set(rec) | set(prev_rec)):
            a, b = prev_rec.get(t), rec.get(t)
            if a is None and b is not None:
                changes.append("  %-5s ENTERED the computed set" % t)
                continue
            if b is None and a is not None:
                changes.append("  %-5s LEFT the computed set -- %s" % (t, blocked.get(t, "not in today's exhibit")))
                continue
            filing = a["anchor"] != b["anchor"] or a.get("roe") != b.get("roe")
            ca, cb = klass(a.get("ratio")), klass(b.get("ratio"))
            dec = decompose(a, b)
            if filing:
                part = (" | ratio %.2f -> %.2f: price %+.2f, filing %+.2f" % (
                    dec["r0"], dec["r1"], dec["price"], dec["filing"])) if dec else ""
                changes.append("  %-5s NEW FILING  anchor %s -> %s   BVPS %s -> %s   12-month ROE %s -> %s%s" % (
                    t, a["anchor"], b["anchor"], fmt(a["bvps"]), fmt(b["bvps"]),
                    fmt(a.get("roe"), "{:.2f}%"), fmt(b.get("roe"), "{:.2f}%"), part))
            if ca != cb:
                why = (" (price %+.2f, filing %+.2f)" % (dec["price"], dec["filing"])) if dec else ""
                changes.append("  %-5s CLASS %s -> %s   ratio %s -> %s%s" % (
                    t, ca, cb, fmt(a.get("ratio")), fmt(b.get("ratio")), why))
            ra = robustness(a)
            if ra["tag"] != rob[t]["tag"]:
                lo, hi = rob[t]["lo"], rob[t]["hi"]
                changes.append("  %-5s ROBUSTNESS %s -> %s   r 7-11%% range %s%s" % (
                    t, ra["tag"], rob[t]["tag"],
                    ("%.2f-%.2f" % (lo, hi)) if lo is not None else "—",
                    ("   (%.2f from the %.2fx line: borderline, expect flicker)" % (
                        rob[t]["edge"], rob[t]["edge_line"])) if rob[t]["borderline"] else ""))
        for t in sorted(set(blocked) - set(prev_blocked)):
            changes.append("  %-5s newly BLOCKED -- %s" % (t, blocked[t]))
    quiet = None
    if changes:
        L.extend(changes)
    else:
        moves = []
        for t, b in rec.items():
            a = prev_rec.get(t)
            if a and a.get("ratio") is not None and b.get("ratio") is not None:
                moves.append((abs(b["ratio"] - a["ratio"]), t, a["ratio"], b["ratio"]))
        if moves:
            mv = max(moves)
            quiet = ("No valuation input changed. Price moves only; no class crossed a line. "
                     "Largest ratio move: %s %.2f -> %.2f" % (mv[1], mv[2], mv[3]))
            L.append("  NO VALUATION INPUT CHANGED. Price moves only; no class crossed a line.")
            L.append("  largest ratio move: %s %.2f -> %.2f" % (mv[1], mv[2], mv[3]))
        else:
            L.append("  nothing to compare.")
    L.append("")

    # ---- 2. decision card --------------------------------------------------
    L.append("-" * 100)
    L.append("2. DECISION CARD -- today's directionals, then held names")
    L.append("-" * 100)
    L.append("  %-5s %-20s %-8s %-24s %-19s %5s %-11s %6s %-11s %-18s %7s %s" % (
        "tk", "held in", "today", "executability", "class (ratio)", "edge", "last-FY",
        "P/E", "r 7-11%", "robust", "g impl", "DP-V2 gate"))
    order = directional + [t for t in sorted(held, key=lambda t: -(rec.get(t, {}).get("ratio") or -1))
                           if t not in directional]
    card = []
    hood = {t: basis.fundamentals(t, date) for t in order if t in basis.cik}
    for t in order:
        s = sig[t]
        f = rec.get(t)
        rb = rob.get(t)
        cls = rb["base"] if rb else "NO READ"
        why = no_read_reason(f, blocked.get(t)) if cls == "NO READ" else ""
        if s["signal"] == "BUY":
            gate = DPV2_GATE[cls]
            flip = dpv2_flip(cls, s["conf"])
            gate_s = "%d%s" % (gate, ("  FLIP " + flip) if flip else "")
        else:
            gate_s = "n/a" if s["signal"] == "SELL" else ""
        cls_s = ("%s (%.2f)" % (cls, f["ratio"])) if f and f.get("ratio") is not None else "NO READ"
        rng = ("%.2f-%.2f" % (rb["lo"], rb["hi"])) if rb and rb["lo"] is not None else "—"
        rtag = rb["tag"] if rb else "—"
        if rb and rtag == "FRAGILE":
            rtag = "FRAGILE r " + "/".join(CLASS_LETTER.get(c, "-") for c in rb["r_classes"])
        if rb and rb["borderline"] and rb["base"] != "NO READ":
            rtag = (rtag + " EDGE") if rtag != "—" else "EDGE"
        fu = hood.get(t) or {}
        eps12 = ((fu.get("growth") or {}).get("eps") or {}).get("ttm")
        px = f.get("px") if f else (s.get("px") or (quotes.get(t) or {}).get("lastPrice"))
        pe12 = (px / eps12) if (px and eps12 and eps12 > 0) else None
        last_s = (rb["lastfy"] if rb and rb["lastfy"] else "—")
        if rb and rb["basis_flip"]:
            last_s = last_s + "*"
        L.append("  %-5s %-20s %-8s %-24s %-19s %5s %-11s %6s %-11s %-18s %7s %s" % (
            t, (s["acct"] if s["held"] else "scout")[:20], "%s %d" % (s["signal"], s["conf"]),
            executability(s)[:24], cls_s, fmt(rb["edge"] if rb else None), last_s[:11],
            fmt(pe12, "{:.1f}"), rng, rtag[:18],
            fmt(f.get("g_impl"), "{:.2f}%") if f else "—", gate_s))
        if cls == "NO READ" and why:
            L.append("  %-5s   NO READ: %s" % ("", why))
        card.append(dict(ticker=t, held=s["held"], account=s["acct"], signal=s["signal"],
                         conf=s["conf"], executability=executability(s), cls=cls,
                         ratio=f.get("ratio") if f else None,
                         edge=rb["edge"] if rb else None, borderline=rb["borderline"] if rb else False,
                         cls_lastfy=rb["lastfy"] if rb else None,
                         ratio_lastfy=f.get("ratio_lastfy") if f else None,
                         pe_ttm=pe12, lo=rb["lo"] if rb else None, hi=rb["hi"] if rb else None,
                         robust=rb["tag"] if rb else "—", robust_detail=rtag,
                         g_impl=f.get("g_impl") if f else None,
                         dpv2_gate=DPV2_GATE[cls] if s["signal"] == "BUY" else None,
                         dpv2_flip=dpv2_flip(cls, s["conf"]) if s["signal"] == "BUY" else None,
                         num_age_q=f.get("num_age") if f else None, no_read_reason=why or None,
                         ttm_roe_pct=f.get("roe") if f else None, ttm_end=f.get("ttm_end") if f else None))
    L.append("")
    L.append("  class: implied ROE at 9% over 12-month ROE. edge: distance to the nearest line")
    L.append("  (under %.2f marks EDGE: a 2-3%% price move can cross it). last-FY: the class on the" % EDGE_BORDERLINE)
    L.append("  basis used before DP-V4, for reference (* where it differs). P/E: price over 12-month")
    L.append("  diluted EPS (SEC). r 7-11%: the ratio across the cost-of-equity grid; ROBUST only")
    L.append("  if all three give one class (FRAGILE r U/N/D lists them, D/N/U = DEMANDING/NEUTRAL/")
    L.append("  UNDEMANDING). DP-V2 gate: the PROPOSED modifier (75/70/65); NOT APPLIED.")
    L.append("")

    # ---- 2b. under the hood (DP-V4) ------------------------------------------
    L.append("  2b. UNDER THE HOOD -- SEC filings, last 12 months against the 12 months before.")
    L.append("      Direction (are earnings growing?), quality (are they cash?), balance sheet.")
    L.append("  %-5s %8s %8s %8s %9s %6s %7s %-18s %-15s %5s %6s %5s %8s  %s" % (
        "tk", "EPS 12m", "EPS y/y", "rev y/y", "lastQ y/y", "OCF/NI", "OCF/rev", "gross mgn (yr ago)",
        "op mgn (yr ago)", "ROIC", "L/E", "L/A", "cash", "note"))

    def mgn(a, b):
        if a is None:
            return "—"
        return "%.0f%% (%s)" % (100 * a, "—" if b is None else "%.0f%%" % (100 * b))
    for c in card:
        t = c["ticker"]
        fu = hood.get(t)
        cav = sector_caveat(t)
        if not fu or fu.get("end") is None:
            reason = ("fund: no company statements"
                      if (quotes.get(t) or {}).get("is_etf") else basis.why_none(t)
                      if not fu else "; ".join(fu.get("notes") or []) or "no 12-month figures")
            L.append("  %-5s %s" % (t, reason))
            c.update(under_hood=None)
            continue
        g, q, b = fu["growth"], fu["quality"], fu["balance"]
        bank = cav.startswith("bank")
        ocfni = None if bank else q.get("ocf_to_ni")
        ocfm = None if bank else q.get("ocf_margin")
        le = None if bank else b.get("liab_to_equity")
        la = None if bank else b.get("liab_to_assets")
        notes = hood_notes(dict(fu, caveat=cav or None))
        L.append("  %-5s %8s %8s %8s %9s %6s %7s %-18s %-15s %5s %6s %5s %8s  %s" % (
            t, fmt(g["eps"]["ttm"]), pct_s(g["eps"]["chg"]), pct_s(g["revenue"]["chg"]),
            pct_s(g["q_eps"]["chg"]), fmt(ocfni), pct0(ocfm),
            mgn(q["gross_margin"], q["gross_margin_year_ago"]),
            mgn(q["op_margin"], q["op_margin_year_ago"]),
            pct0(roic_value(dict(fu, caveat=cav or None))), fmt(le), pct0(la),
            ("%.1fB" % (b["cash"] / 1e9)) if b.get("cash") is not None else "—",
            "; ".join(notes)[:120]))
        c.update(under_hood=dict(fu, caveat=cav or None))
    L.append("      y/y: change against the same window a year earlier, signed by direction;")
    L.append("      beyond +/-300% the base is too small to read as a rate. OCF/NI: operating cash")
    L.append("      flow over net income (below 1: profits not yet cash). OCF/rev: operating cash")
    L.append("      flow over revenue, defined through a loss. L/E: liabilities / equity. L/A:")
    L.append("      liabilities / assets, defined when equity is negative (above 100%: a deficit).")
    L.append("      ROIC (DP-V6): 12-month operating profit after tax / average invested capital")
    L.append("      (debt + equity incl. minority interests - cash). Tax: the filer's own rate, else")
    L.append("      its 3-year rate, else 25.5%; an operating loss gets no shield. Defined when equity")
    L.append("      is negative; a bank's is not shown. The note says when a read is not standard.")
    L.append("")

    # ---- 2c. expectations (DP-V4) ---------------------------------------------
    L.append("  2c. EXPECTATIONS -- the next report and Finnhub's ADJUSTED (non-GAAP) consensus,")
    L.append("      on their own basis. 'class if met' scales the 12-month ROE by the change the")
    L.append("      consensus implies, book unchanged: APPROXIMATE, context only, never the class.")
    has_earn = bool(ctx.get("earn"))
    if not has_earn and not ctx.get("ttm"):
        L.append("      not available for %s (value.bat did not write the calendar)." % date)
    L.append("  %-5s %-11s %8s %10s %9s %9s %-16s %-14s %s" % (
        "tk", "next report", "sessions", "consensus", "beat (4)", "adj 12m", "if met",
        "class if met", "note"))
    for c in card:
        t = c["ticker"]
        _e, _roe, nd, ns, note = _ctx_cells(ctx, t)
        ee = (((ctx.get("earn") or {}).get(t)) or {}).get("expectations") or {}
        chg = ee.get("chg_if_met")
        cim = None
        if chg is not None and c.get("ratio") is not None and (1.0 + chg) > 0:
            cim = c["ratio"] / (1.0 + chg)
        full_note = note
        met_s = ("%s (%s)" % (fmt(ee.get("adj_ttm_if_met")), pct_s(chg))
                 if ee.get("adj_ttm_if_met") is not None else (ee.get("why") or "—"))
        L.append("  %-5s %-11s %8s %10s %9s %9s %-16s %-14s %s" % (
            t, nd or "—", "" if ns is None else ns, fmt(ee.get("eps_estimate")),
            ("%d of %d" % (ee["beats"], ee["reported"])) if ee.get("beats") is not None else "—",
            fmt(ee.get("adj_ttm")), met_s[:16],
            ("%s (%.2f)" % (klass(cim), cim)) if cim is not None else "—", full_note[:140]))
        c.update(next_report=nd, sessions_to_report=ns, context_note=full_note,
                 ttm_eps=((c.get("under_hood") or {}).get("growth") or {}).get("eps", {}).get("ttm"),
                 expectations=dict(ee, cls_if_met=klass(cim) if cim is not None else None,
                                   ratio_if_met=cim) if ee else None)
    L.append("")

    # ---- 3. robustness, every classed name ---------------------------------
    L.append("-" * 100)
    L.append("3. ROBUSTNESS -- every classed name: the cost-of-equity grid on the 12-month basis")
    L.append("-" * 100)
    L.append("  %-5s %-12s %7s %-20s | %-12s %-12s %-12s | %6s %-8s" % (
        "tk", "class @9%", "ratio", "last-FY (ratio)", "r=7%", "r=9%", "r=11%", "edge", "tag"))
    for t in sorted(read, key=lambda t: -(rec[t].get("ratio") or 0)):
        rb = rob[t]
        L.append("  %-5s %-12s %7s %-20s | %-12s %-12s %-12s | %6s %-8s" % (
            t, rb["base"], fmt(rec[t].get("ratio")),
            ("%s (%s)" % (rb["lastfy"], fmt(rb["lastfy_ratio"])))[:20],
            *rb["r_classes"], fmt(rb["edge"]), rb["tag"] + (" EDGE" if rb["borderline"] else "")))
    n_rob = sum(1 for t in read if rob[t]["tag"] == "ROBUST")
    n_edge = sum(1 for t in read if rob[t]["borderline"])
    n_basis = sum(1 for t in read if rob[t]["basis_flip"])
    L.append("  ROBUST %d of %d.  EDGE (within %.2f of a line) %d.  Class differs from last-FY: %d." % (
        n_rob, len(read), EDGE_BORDERLINE, n_edge, n_basis))
    L.append("")

    # ---- 4. DP-V2 counterfactual flip ledger -------------------------------
    L.append("-" * 100)
    L.append("4. DP-V2 COUNTERFACTUAL FLIP LEDGER -- COUNTS ONLY. No outcome is read here.")
    L.append("-" * 100)
    cls_ttm, cls_old = {}, {}
    for d, p in exhibits.items():
        if DPV2_SINCE <= d <= date:
            r_, _ = parse_exhibit(p)
            # the BUY on d was decided that morning: only filings public BEFORE d
            cls_ttm[d] = {t: klass(f.get("ratio")) for t, f in rebase(r_, basis, d, strict=True).items()}
            if d < DPV4_ON:
                cls_old[d] = {t: klass(f.get("ratio")) for t, f in r_.items()}
    flips, unscored, closed, closed_un = [], [], [], []
    if tracker:
        rows = [r for r in load_buy_rows(tracker, DPV2_SINCE) if r[0] <= date]
        flips, unscored = dpv2_ledger(rows, cls_ttm, label=_label_ttm)
        closed, closed_un = dpv2_ledger([r for r in rows if r[0] < DPV4_ON], cls_old)
    L.append("  ON THE 12-MONTH BASIS (DP-V4; the count). Rows before %s are REBUILT point in time:" % DPV4_ON)
    L.append("  only filings public before that session, and the exhibit written that day.")
    for d, tk, c, cls, f, src in flips:
        L.append("  %s  %-5s BUY %3.0f  %-11s %-14s %s" % (d, tk, c, cls, f, src))
    if not flips:
        L.append("  no flip rows.")
    for d, tk, c, hint, pd_ in unscored:
        L.append("  %s  %-5s BUY %3.0f  UNSCORED -- no exhibit that session; %s on %s would flip it. "
                 "Not counted [?]" % (d, tk, c, hint, pd_))
    try:
        import trading_calendar as cal
        left = len(cal.sessions_between(date + dt.timedelta(days=1), DPV2_STOP))
    except Exception:                                              # noqa: BLE001
        left = None
    tks = sorted({f[1] for f in flips})
    L.append("  flip rows %d of %d needed   tickers %d of %d (%s)   stop %s%s" % (
        len(flips), DPV2_MIN_ROWS, len(tks), DPV2_MIN_TICKERS, ", ".join(tks) or "none",
        DPV2_STOP, ("   sessions left %d" % left) if left is not None else ""))
    L.append("  If the n condition is unmet on %s, DP-V2 is WITHDRAWN, NOT EXTENDED (memo 6)." % DPV2_STOP)
    closed_tks = sorted({f[1] for f in closed})
    L.append("")
    L.append("  CLOSED -- the last-FY ledger as it ran %s .. %s (DP-V4). On the record; never" % (
        DPV2_SINCE, DPV4_ON - dt.timedelta(days=1)))
    L.append("  pooled with the count above:")
    for d, tk, c, cls, f, src in closed:
        L.append("  %s  %-5s BUY %3.0f  %-11s %-14s %s" % (d, tk, c, cls, f, src))
    for d, tk, c, hint, pd_ in closed_un:
        L.append("  %s  %-5s BUY %3.0f  UNSCORED (prior %s on %s)" % (d, tk, c, hint, pd_))
    L.append("  closed: %d row(s), %d ticker(s) (%s)" % (len(closed), len(closed_tks),
                                                       ", ".join(closed_tks) or "none"))
    L.append("")

    # ---- 5. book-free reads ------------------------------------------------
    L.append("-" * 100)
    L.append("5. BOOK-FREE READS -- context only, for names the class cannot read")
    L.append("-" * 100)
    L.append("  P/E band: the class's own lines restated as P/E (about %.0f and %.0f, section 6)," % (
        DEMANDING_AT / ((R_BASE - G_BASE) / 100 / (1 + G_BASE / 100)),
        NEUTRAL_AT / ((R_BASE - G_BASE) / 100 / (1 + G_BASE / 100))))
    L.append("  on 12-month EPS (SEC, else Schwab's own 12-month P/E). Dividend-implied g: Gordon at")
    L.append("  r = 9%, g = (r - y)/(1 + y). It ignores buybacks, so it UNDERSTATES for buyback-heavy names.")
    pe_hi = DEMANDING_AT / ((R_BASE - G_BASE) / (100 + G_BASE))
    pe_lo = NEUTRAL_AT / ((R_BASE - G_BASE) / (100 + G_BASE))
    bf = []
    for t in sorted(set(held) | set(directional)):
        if t in rec and rob[t]["base"] != "NO READ":
            continue
        q = quotes.get(t, {})
        px, y = sig[t].get("px") or q.get("lastPrice"), q.get("divYield")
        fu = hood.get(t) if t in hood else basis.fundamentals(t, date) if t in basis.cik else None
        eps12 = (((fu or {}).get("growth") or {}).get("eps") or {}).get("ttm")
        pe, src = None, ""
        if eps12 and eps12 > 0 and px:
            pe, src = px / eps12, "SEC 12m"
        elif q.get("peRatio") and q["peRatio"] > 0:
            pe, src = q["peRatio"], "Schwab 12m"
        band = None if pe is None else ("DEMANDING-equiv" if pe >= pe_hi else
                                        ("NEUTRAL-equiv" if pe >= pe_lo else "UNDEMANDING-equiv"))
        gdiv = dividend_implied_g(y)
        etf = q.get("is_etf") or q.get("asset_type") not in (None, "EQUITY")
        note = "fund: distributions, not earnings" if etf else src
        bf.append((t, pe, band, y, gdiv, note))
        L.append("  %-5s P/E %7s  %-18s divY %6s  dividend-implied g %7s  %s" % (
            t, fmt(pe, "{:.1f}"), band or "no positive eps", fmt(y, "{:.2f}%"),
            fmt(gdiv, "{:.2f}%"), note))
    L.append("")

    # ---- 6. identity check -------------------------------------------------
    L.append("-" * 100)
    L.append("6. IDENTITY CHECK -- implied_return's own ratio (its last-FY basis) as a P/E:")
    L.append("   (P/EPS)(r-g)/(1+g) + g/((1+g)ROE). It checks implied_return's arithmetic, not the class.")
    L.append("-" * 100)
    worst, breach, checked = 0.0, [], 0
    for t, f in rec_raw.items():
        if not f.get("roe") or f["roe"] <= 0 or f.get("ratio") is None:
            continue
        recon = ratio_from(f["px"], f["bvps"], f["roe"])
        if recon is None:
            continue
        checked += 1
        diff = abs(recon - f["ratio"])
        bound = f["ratio"] * (0.005 / f["px"] + 0.005 / f["bvps"] + 0.005 / f["roe"]) + 0.005
        worst = max(worst, diff)
        if diff > bound:
            breach.append("%s %.3f vs %.2f (bound %.3f)" % (t, recon, f["ratio"], bound))
    k = (R_BASE - G_BASE) / (100.0 + G_BASE)
    L.append("  %d names reconstructed; max |reconstructed - printed| = %.3f; outside the rounding "
             "bound: %s" % (checked, worst, ", ".join(breach) or "none"))
    L.append("  so a class is %.4f x P/E + %.4f/ROE: DEMANDING ~ P/E >= %.1f, UNDEMANDING ~ < %.1f" % (
        k, G_BASE / (100.0 + G_BASE), DEMANDING_AT / k, NEUTRAL_AT / k))
    L.append("  on whichever EPS the ROE carries: since DP-V4, the 12-month EPS.")
    if breach:
        findings.append("identity check FAILED for %s: implied_return's class no longer reconstructs "
                        "as a P/E within rounding -- its formula or inputs changed; re-read section 1"
                        % ", ".join(b.split()[0] for b in breach))
    unparsed = [t for t in sig if t not in rec and t not in blocked]
    L.append("")
    L.append("=" * 100)
    L.append("FINDINGS: %d" % len(findings))
    for x in findings:
        L.append("  - " + x)
    L.append("  not in the valuation set (dark, fund, scope or negative ROE): %s" % (", ".join(sorted(unparsed)) or "none"))
    L.append("=" * 100)
    L.append("Counterfactual only. No gate modified. No valuation asserted, no price target implied.")
    L.append("Advisory-only - no orders - not financial advice.")

    # ASCII only on the way out: the runners redirect stdout through the Windows
    # ANSI code page, and a stray non-ASCII dash should never cost an exhibit.
    L = [line.replace("—", "-") for line in L]
    snapshot = dict(date=str(date), exhibit=os.path.basename(exhibits[date]),
                    log=os.path.basename(log_path), coverage=coverage, quiet=quiet,
                    basis=dict(name="12-month ROE (DP-V4)", since=str(DPV4_ON), line=basis_line,
                               differ=[dict(ticker=t, lastfy=rob[t]["lastfy"], now=rob[t]["base"])
                                       for t in sorted(basis_diff)]),
                    identity_outside_bound=breach,
                    identity=dict(checked=checked, max_diff=worst, k=k,
                                  pe_demanding=DEMANDING_AT / k, pe_undemanding=NEUTRAL_AT / k),
                    previous=str(prev_date) if prev_date else None, card=card,
                    robustness={t: dict(cls=rob[t]["base"], ratio=rec[t].get("ratio"),
                                        lastfy=rob[t]["lastfy"], lastfy_ratio=rob[t]["lastfy_ratio"],
                                        tag=rob[t]["tag"], edge=rob[t]["edge"],
                                        borderline=rob[t]["borderline"],
                                        lo=rob[t]["lo"], hi=rob[t]["hi"],
                                        r_classes=rob[t]["r_classes"], ttm_roe_pct=rec[t].get("roe"),
                                        ttm_why=rec[t].get("ttm_why")) for t in rec},
                    changes=[c.strip() for c in changes],
                    dpv2=dict(basis="12-month (DP-V4)", flip_rows=len(flips), tickers=tks,
                              needed_rows=DPV2_MIN_ROWS, needed_tickers=DPV2_MIN_TICKERS,
                              stop=str(DPV2_STOP), sessions_left=left,
                              unscored=[dict(date=str(u[0]), ticker=u[1], conf=u[2], hint=u[3])
                                        for u in unscored],
                              rows=[dict(date=str(d), ticker=tk, conf=c, cls=cl, flip=f, source=s)
                                    for d, tk, c, cl, f, s in flips],
                              closed=dict(basis="last-FY (closed by DP-V4)", rows=[
                                  dict(date=str(d), ticker=tk, conf=c, cls=cl, flip=f, source=s)
                                  for d, tk, c, cl, f, s in closed], tickers=closed_tks)),
                    book_free=[dict(ticker=t, pe=pe, band=b, div_yield=y, g_div=g, note=n)
                               for t, pe, b, y, g, n in bf],
                    identity_max_diff=worst, findings=findings)
    return "\n".join(L), snapshot, findings


# --------------------------------------------------------------------------- #
# Embed into the GENERATED report copy (value.bat step 5b, added 2026-10-14)
# --------------------------------------------------------------------------- #
#
# The report copy (tracker_with_registry.xlsx) is the file read as the final
# word, so the card belongs in it. Four rules, each with the failure it stops:
#
#   1. NEVER WHILE EXCEL HAS IT OPEN. Excel holds the file with a deny-write
#      share lock and keeps an owner file "~$<name>" beside it. Writing under
#      it either fails half-way or is silently overwritten when Excel next
#      saves. Both are checked; a stale owner file (Excel closed) is refused
#      too, with the instruction to delete it -- the 2026-08-04 one sat in the
#      root for ten weeks.
#   2. NEVER INTO THE MASTER. The master is manual entry only.
#   3. ONLY INTO TODAY'S COPY. build_near_miss_registry --emit-master-copy
#      rebuilds the copy from the master on every review run (and drops every
#      embedded sheet). If the copy's latest cohort is not the card's date, the
#      review has not run today: embedding would put today's card into
#      yesterday's report.
#   4. ATOMIC. Saved to a temporary file beside the copy, then swapped in with
#      os.replace, so a failure never leaves a half-written report copy.
#
# openpyxl does not keep Excel's cached formula results when it saves. The
# workbook is therefore flagged fullCalcOnLoad, and the Excel refresh (daily
# workflow step 6) belongs AFTER value.bat, not before it.

VC_SHEET = "Value Card"
VC_AFTER = ("Monitor Strip", "Actionable BUYs", "Near-Miss BUYs", "Signals")
MASTER_NAME = "claude_equity_bot_tracker.xlsx"


def _probe_write(path):
    """Open for writing and close at once. Excel's open denies other writers on
    Windows, so this raises PermissionError while the workbook is open."""
    fh = open(path, "r+b")
    fh.close()


def excel_holds(path):
    """(reason, detail) if Excel has, or appears to have, the workbook open;
    (None, None) when it is safe to write."""
    folder, base = os.path.split(os.path.abspath(path))
    owner = os.path.join(folder, "~$" + base)
    try:
        _probe_write(path)
    except PermissionError:
        return "IN USE", "Excel (or another program) has %s open. Close it, then re-run value.bat" % base
    if os.path.exists(owner):
        return "OWNER FILE", ("%s exists. If Excel is closed it is stale: delete it, then re-run"
                              % os.path.basename(owner))
    return None, None


def latest_cohort(wb):
    """Newest date in the Signals sheet's date column."""
    from build_actionable_registry import SIGNALS_SHEET, FIRST_DATA_ROW, COL, as_date
    ws = wb[SIGNALS_SHEET]
    best, blank = None, 0
    for (v,) in ws.iter_rows(min_row=FIRST_DATA_ROW, min_col=COL["date"], max_col=COL["date"],
                             values_only=True):
        d = as_date(v)
        if d is None:
            blank += 1
            if blank >= 25:
                break
            continue
        blank = 0
        best = d if best is None or d > best else best
    return best


def _write_sheet(ws, snap):
    from openpyxl.styles import Alignment, Font, PatternFill
    f_title = Font(bold=True, size=13)
    f_sec = Font(bold=True, size=11)
    f_hdr = Font(bold=True, size=10)
    f_note = Font(italic=True, size=9, color="FF6B6B6B")
    fill_hdr = PatternFill("solid", fgColor="FFE7E6E1")
    fill_rob = PatternFill("solid", fgColor="FFDDEFD9")
    fill_frag = PatternFill("solid", fgColor="FFF6E7C8")
    top = Alignment(vertical="top", wrap_text=False)
    r = [1]

    def put(values, font=None, fill=None):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r[0], c, v)
            cell.alignment = top
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
        r[0] += 1

    def gap():
        r[0] += 1

    d = snap["date"]
    put(["Value Card -- %s" % d], f_title)
    put(["as of %s | built by value_card.py from %s, %s, quote_dump.json and the master | "
         "COUNTERFACTUAL ONLY: modifies no gate, reads no outcome | advisory-only, no orders"
         % (d, snap["exhibit"], snap.get("log", "the signals log"))], f_note)
    gap()
    cv = snap["coverage"]
    put(["Coverage"], f_sec)
    put(["watchlist %d, classed %d, ROBUST %d | held %d, classed %d, ROBUST %d (%s) | "
         "today's directionals %d, classed %d (%s)" % (
             cv["watchlist"], cv["classed"], cv["robust"], cv["held"], cv["held_classed"],
             len(cv["held_robust"]), ", ".join(cv["held_robust"]) or "none",
             cv["directional"], len(cv["directional_classed"]),
             ", ".join(cv["directional_classed"]) or "none")])
    gap()
    if snap.get("basis"):
        put([snap["basis"]["line"]], f_sec)
        gap()
    put(["What changed since %s" % (snap.get("previous") or "(no earlier exhibit)")], f_sec)
    for line in (snap["changes"] or [snap.get("quiet") or "Nothing to compare."]):
        put([line])
    gap()
    put(["Decision card -- today's directionals, then held names (class on 12-month ROE, DP-V4)"], f_sec)
    hdr = ["Ticker", "Held in", "Today", "Conf", "Executability", "Class (12-month)", "Ratio",
           "Edge to line", "Last-FY class", "Last-FY ratio", "P/E (12-month)", "Range low (r 7-11%)",
           "Range high", "Robust", "Why fragile", "Implied g %", "DP-V2 gate (not applied)",
           "Would flip", "NO READ reason", "12-month ROE %", "Next report", "Sessions to report",
           "Context note"]
    put(hdr, f_hdr, fill_hdr)
    for c in snap["card"]:
        detail = (c.get("robust_detail") or "").replace("FRAGILE", "").strip()
        tag = c["robust"] if c["robust"] != "—" else ""
        put([c["ticker"], c["account"] if c["held"] else "scout", c["signal"], c["conf"],
             c["executability"].replace("—", ""), c["cls"], c["ratio"], c.get("edge"),
             c.get("cls_lastfy") or "", c.get("ratio_lastfy"), c.get("pe_ttm"),
             c["lo"], c["hi"], tag, detail, c["g_impl"], c.get("dpv2_gate"),
             c.get("dpv2_flip") or "", c.get("no_read_reason") or "",
             c.get("ttm_roe_pct"), c.get("next_report") or "",
             c.get("sessions_to_report"), c.get("context_note") or ""],
            fill=fill_rob if tag == "ROBUST" else (fill_frag if tag == "FRAGILE" else None))
    put(["Class: implied ROE at 9%% over 12-month ROE (SEC, point in time). Edge: distance to the "
         "nearest line (under %.2f, a 2-3%% price move can cross it). Range: r = 7/9/11%%; ROBUST "
         "only if all three give one class. Last-FY: the basis before DP-V4, for reference."
         % EDGE_BORDERLINE], f_note)
    gap()
    put(["Under the hood -- SEC filings, last 12 months against the 12 months before"], f_sec)
    put(list(HOOD_SHEET_COLUMNS), f_hdr, fill_hdr)
    for c in snap["card"]:
        u = c.get("under_hood")
        if not u:
            continue
        put(hood_sheet_row(c["ticker"], u))
    gap()
    put(["Expectations -- Finnhub's ADJUSTED (non-GAAP) consensus, on its own basis; "
         "'class if met' is approximate and never the class"], f_sec)
    put(["Ticker", "Next report", "Sessions", "Consensus EPS (adj.)", "Beat estimate (of last 4)",
         "Adj. 12-month EPS", "Adj. 12-month EPS if met", "Change if met %", "Class if met (approx.)",
         "Ratio if met", "Note"], f_hdr, fill_hdr)
    for c in snap["card"]:
        e = c.get("expectations") or {}
        put([c["ticker"], c.get("next_report") or "", c.get("sessions_to_report"),
             e.get("eps_estimate"),
             ("%d of %d" % (e["beats"], e["reported"])) if e.get("beats") is not None else "",
             e.get("adj_ttm"), e.get("adj_ttm_if_met"), _p100(e.get("chg_if_met")),
             e.get("cls_if_met") or "", e.get("ratio_if_met"),
             e.get("why") or e.get("history_why") or ""])
    gap()
    put(["Robustness -- every classed name: the cost-of-equity grid on the 12-month basis"], f_sec)
    put(["Ticker", "Class @9%", "Ratio", "Last-FY class", "Last-FY ratio", "r = 7%", "r = 9%",
         "r = 11%", "Edge", "Tag"], f_hdr, fill_hdr)
    rb = snap["robustness"]
    for t in sorted((t for t in rb if rb[t]["cls"] != "NO READ"),
                    key=lambda t: -(rb[t]["ratio"] or 0)):
        x = rb[t]
        put([t, x["cls"], x["ratio"], x.get("lastfy") or "", x.get("lastfy_ratio"),
             *x.get("r_classes", ["", "", ""]), x.get("edge"),
             x["tag"] + (" EDGE" if x.get("borderline") else "")],
            fill=fill_rob if x["tag"] == "ROBUST" else fill_frag)
    gap()
    dp = snap["dpv2"]
    put(["DP-V2 counterfactual flip ledger -- COUNTS ONLY, no outcome read (%s)"
         % dp.get("basis", "last-FY")], f_sec)
    put(["flip rows %d of %d needed | tickers %d of %d (%s) | stop %s | sessions left %s | "
         "withdrawn, not extended, if unmet on the stop date" % (
             dp["flip_rows"], dp["needed_rows"], len(dp["tickers"]), dp["needed_tickers"],
             ", ".join(dp["tickers"]) or "none", dp["stop"], dp.get("sessions_left"))])
    put(["Date", "Ticker", "BUY conf", "Class that day", "Flip", "Source"], f_hdr, fill_hdr)
    for x in dp["rows"]:
        put([x["date"], x["ticker"], x["conf"], x["cls"], x["flip"], x["source"]])
    for x in dp.get("unscored", []):
        put([x["date"], x["ticker"], x["conf"], "UNSCORED (prior: %s)" % x["hint"],
             "not counted", "no exhibit that session"])
    cl = dp.get("closed")
    if cl:
        put(["CLOSED by DP-V4 -- the last-FY ledger (%d rows; %s). On the record, never pooled "
             "with the count above." % (len(cl["rows"]), ", ".join(cl["tickers"]) or "none")], f_note)
        for x in cl["rows"]:
            put([x["date"], x["ticker"], x["conf"], x["cls"], x["flip"], x["source"] + " (closed)"])
    gap()
    put(["Book-free reads -- context only"], f_sec)
    put(["Ticker", "P/E (12-month EPS)", "P/E band", "Div yield %",
         "Dividend-implied g %", "Note"], f_hdr, fill_hdr)
    for x in snap["book_free"]:
        put([x["ticker"], x["pe"], x["band"] or "no positive eps", x["div_yield"], x["g_div"],
             x["note"]])
    gap()
    idn = snap.get("identity", {})
    put(["Identity check"], f_sec)
    put(["implied_return's ratio = %.4f x P/E + %.4f/ROE (DEMANDING ~ P/E >= %.1f, UNDEMANDING ~ < %.1f) | "
         "%d names reconstructed, max diff %.3f | outside rounding bound: %s" % (
             idn.get("k", 0), G_BASE / (100.0 + G_BASE), idn.get("pe_demanding", 0),
             idn.get("pe_undemanding", 0), idn.get("checked", 0), idn.get("max_diff", 0),
             ", ".join(snap.get("identity_outside_bound") or []) or "none")])
    if snap.get("findings"):
        put(["FINDINGS: " + " | ".join(snap["findings"])], f_sec)
    # number formats and widths
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    widths = [8, 22, 9, 8, 26, 14, 8, 10, 13, 10, 10, 10, 10, 9, 14, 11, 12, 15, 40, 10, 12, 10, 40]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A3"


def embed(snap, copy_path, master_path=None):
    """Write/refresh the Value Card sheet in the report copy. Returns
    (ok, message). Never raises for an expected refusal."""
    import openpyxl
    base = os.path.basename(copy_path).lower()
    if base == MASTER_NAME or (master_path and os.path.abspath(copy_path) == os.path.abspath(master_path)):
        return False, "REFUSED: %s is the MASTER. The card is embedded into the report copy only." % copy_path
    if not os.path.exists(copy_path):
        return False, "REFUSED: %s not found. Run run_daily_review.bat first." % copy_path
    why, detail = excel_holds(copy_path)
    if why:
        return False, "REFUSED (%s): %s." % (why, detail)
    wb = openpyxl.load_workbook(copy_path)            # data_only=False: keep formulas
    newest = latest_cohort(wb)
    if str(newest) != snap["date"]:
        wb.close()
        return False, ("REFUSED: the report copy's latest cohort is %s, the card's is %s -- "
                       "run_daily_review.bat has not rebuilt it for today." % (newest, snap["date"]))
    if VC_SHEET in wb.sheetnames:
        del wb[VC_SHEET]                              # replace, never stack
    ws = wb.create_sheet(VC_SHEET)
    _write_sheet(ws, snap)
    for after in VC_AFTER:
        if after in wb.sheetnames and after != VC_SHEET:
            target = wb.sheetnames.index(after) + 1
            wb.move_sheet(VC_SHEET, offset=target - wb.sheetnames.index(VC_SHEET))
            break
    try:
        wb.calculation.fullCalcOnLoad = True          # cached results are gone; recalc on open
    except AttributeError:
        pass
    tmp = copy_path + ".vc-tmp"
    try:
        wb.save(tmp)
        why, detail = excel_holds(copy_path)          # re-check: Excel may have opened it meanwhile
        if why:
            return False, "REFUSED (%s) at save time: %s. Nothing was changed." % (why, detail.rstrip("."))
        os.replace(tmp, copy_path)
    except PermissionError as e:
        return False, "REFUSED: could not replace %s (%s). Nothing was changed." % (copy_path, e)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return True, "EMBEDDED: sheet '%s' (%d card rows) written into %s for %s." % (
        VC_SHEET, len(snap["card"]), os.path.basename(copy_path), snap["date"])


# --------------------------------------------------------------------------- #
# Self-test (positive and negative cases)
# --------------------------------------------------------------------------- #

def _selftest():
    import tempfile
    fails = []

    def check(name, ok):
        print("  [%s] %s" % ("PASS" if ok else "FAIL", name))
        if not ok:
            fails.append(name)

    # thresholds and flips
    check("2.00x is DEMANDING", klass(2.00) == "DEMANDING")
    check("1.9999x is NEUTRAL", klass(1.9999) == "NEUTRAL")
    check("1.50x is NEUTRAL", klass(1.50) == "NEUTRAL")
    check("1.4999x is UNDEMANDING", klass(1.4999) == "UNDEMANDING")
    check("no ratio is NO READ", klass(None) == "NO READ")
    check("DEMANDING BUY 74 flips ADMIT->REJECT", dpv2_flip("DEMANDING", 74) == "ADMIT->REJECT")
    check("DEMANDING BUY 75 does not flip", dpv2_flip("DEMANDING", 75) is None)
    check("UNDEMANDING BUY 65 flips REJECT->ADMIT", dpv2_flip("UNDEMANDING", 65) == "REJECT->ADMIT")
    check("UNDEMANDING BUY 64 does not flip", dpv2_flip("UNDEMANDING", 64) is None)
    check("NEUTRAL BUY 72 does not flip", dpv2_flip("NEUTRAL", 72) is None)
    check("NO READ BUY 70 does not flip", dpv2_flip("NO READ", 70) is None)

    # identity: NVDA 2026-10-14 as printed by implied_return
    nv = dict(px=223.70, bvps=9.44, roe=51.92)
    check("identity reproduces NVDA 2.72x", abs(ratio_from(**nv) - 2.72) < 0.02)
    eps = nv["roe"] / 100 * nv["bvps"]
    k = (R_BASE - G_BASE) / (100 + G_BASE)
    check("ratio == k*P/E + g/((1+g)ROE) exactly",
          abs(ratio_from(**nv) - (k * nv["px"] / eps + (0.03 / 1.03) / 0.5192)) < 1e-9)
    check("BVPS cancels: doubling BVPS and halving ROE leaves the P/E term unchanged",
          abs(k * nv["px"] / eps - k * nv["px"] / ((nv["roe"] / 2 / 100) * (nv["bvps"] * 2))) < 1e-9)
    check("NEGATIVE: non-positive ROE gives no ratio", ratio_from(100, 10, -5) is None)

    # decomposition sums to the total
    a = dict(px=210.57, bvps=8.01, roe=61.17)
    b = dict(px=218.16, bvps=9.44, roe=51.92)
    dec = decompose(a, b)
    check("price + filing == total move", abs(dec["price"] + dec["filing"] - (dec["r1"] - dec["r0"])) < 1e-12)

    # robustness on a synthetic record (DP-V4: the r grid on 12-month ROE)
    f = dict(ratio=140.99 / 84.2, roe=84.2, i7=94.97, i9=140.99, i11=187.02, ratio_lastfy=2.72)
    rb = robustness(f)
    check("DP-V4: an NVDA-like record on 12-month ROE is NEUTRAL, FRAGILE across r, and "
          "differs from its last-FY class", rb["base"] == "NEUTRAL" and rb["tag"] == "FRAGILE"
          and rb["r_flip"] and rb["basis_flip"] and rb["lastfy"] == "DEMANDING")
    f2 = dict(ratio=6.03, roe=10.76, i7=44.21, i9=64.87, i11=85.52)
    check("PWR-like record is ROBUST", robustness(f2)["tag"] == "ROBUST")
    f3 = dict(ratio=52.97 / 35.9, roe=35.9, i7=36.28, i9=52.97, i11=69.65)
    r3 = robustness(f3)
    check("a ratio within 0.05 of a line is EDGE (FTI 1.48 against 1.50)",
          r3["borderline"] and r3["edge_line"] == NEUTRAL_AT and r3["base"] == "UNDEMANDING")
    check("(III) on 12-month ROE reproduces implied_return's g for NVDA (6.95%)",
          abs(g_implied_pct(51.92, 23.70) - 6.95) < 0.01 and g_implied_pct(10.0, -1.0) is None)

    # Gordon solved for g
    g = dividend_implied_g(3.0)
    check("dividend-implied g at y=3% is (9-3)/1.03", abs(g - 6.0 / 1.03) < 1e-9)
    check("NEGATIVE: no yield gives no g", dividend_implied_g(0) is None)

    # parsing real-format lines, then the end-to-end refusal paths
    tmp = tempfile.mkdtemp(prefix="value_card_selftest_")
    ex = "\n".join([
        "-" * 100, "1. INPUTS -- ROE is trailing eps over book value per share", "-" * 100,
        "    tk      anchor      BVPS  px(master)   drift      ROE     M/B  divY%  anc.age num.age  ROE basis",
        "  NVDA  2026-07-26      9.44      223.70  -1.01%   51.92%   23.70   0.44     0.9Q    2.9Q  Schwab last-FY",
        "   FTI  2026-06-30      8.12       69.76  -1.32%   28.33%    8.59   0.29     1.3Q    3.3Q  Schwab last-FY",
        "", "-" * 100, "1a. THE TWO ROE ROUTES", "-" * 100,
        "  NVDA   Q2      51.92%     104.27%   +52.34  EDGAR quarterly x4",
        "", "-" * 100, "4. (II) WHAT ROE MUST THE MARKET BELIEVE?   <-- the headline read", "-" * 100,
        "  NVDA   23.70     51.92%     94.97%    140.99%    187.02%                   2.72x",
        "   FTI    8.59     28.33%     36.28%     52.97%     69.65%                   1.87x",
        "", "-" * 100, "4b. (III) WHAT GROWTH MUST THE MARKET BELIEVE?", "-" * 100,
        "  NVDA   23.70   0.9797        5.13%          6.95%  g (III)",
        "", "-" * 100, "5. WHY", "-" * 100,
        "         PM NOT COMPUTED -- book equity is non-positive (-8.58B); M/B undefined", ""])
    p = os.path.join(tmp, "implied_return_2026-10-14.md")
    with open(p, "w", encoding="utf-8") as fh:
        fh.write(ex)
    rec, bl = parse_exhibit(p)
    check("parses section 1 with the vintage columns", rec.get("NVDA", {}).get("num_age") == 2.9)
    check("parses the EDGAR route", rec["NVDA"].get("roe_edgar") == 104.27)
    check("parses the class ratio and g", rec["NVDA"].get("ratio") == 2.72 and rec["NVDA"].get("g_impl") == 6.95)
    check("parses a blocked name", "PM" in bl)
    log = os.path.join(tmp, "signals_2026-10-14.log")
    with open(log, "w", encoding="utf-8") as fh:
        fh.write("x | [PHASE 3-C DRY RUN] [Roth IRA,Individual] [HELD] NVDA | Signal: BUY | "
                 "Confidence: 74% | Risk: APPROVED | Reason: x PER-ACCOUNT: Roth IRA=APPROVED; "
                 "Individual=REJECTED. | H1[]\n")
        fh.write("x | [PHASE 3-C DRY RUN] [Acct A,Acct B] [HELD] FTI | Signal: BUY | "
                 "Confidence: 68% | Risk: REJECTED | Reason: x PER-ACCOUNT: Acct A=REJECTED; "
                 "Acct B=REJECTED. | H1[]\n")
    s = parse_log(log)
    check("parses held, accounts and per-account", s["NVDA"]["held"] and
          executability(s["NVDA"]) == "PARTIAL (Roth IRA only)")
    # ---- DP-V4: a two-filer slim cache, point in time ----------------------
    bare = tempfile.mkdtemp(prefix="value_card_bare_")       # no cache at all
    with open(os.path.join(bare, "implied_return_2026-10-14.md"), "w", encoding="utf-8") as fh:
        fh.write(ex)

    def _qs(ends, vals, lag=30):
        out, prev = [], None
        for e, v in zip(ends, vals):
            s0 = (dt.date.fromisoformat(prev) + dt.timedelta(days=1)).isoformat() if prev else \
                (dt.date.fromisoformat(e) - dt.timedelta(days=90)).isoformat()
            out.append({"start": s0, "end": e, "val": v, "form": "10-Q", "accn": "a" + e,
                        "filed": (dt.date.fromisoformat(e) + dt.timedelta(days=lag)).isoformat()})
            prev = e
        return out
    nv_ends = ["2024-10-27", "2025-01-26", "2025-04-27", "2025-07-27",
               "2025-10-26", "2026-01-25", "2026-04-26", "2026-07-26"]
    ft_ends = ["2024-09-30", "2024-12-31", "2025-03-31", "2025-06-30",
               "2025-09-30", "2025-12-31", "2026-03-31", "2026-06-30"]
    nvda = {"slim_version": 3, "facts": {"us-gaap": {
        "NetIncomeLoss": {"units": {"USD": _qs(nv_ends, [19e9, 22e9, 19e9, 26e9,
                                                         31.91e9, 42.96e9, 58.32e9, 59.69e9])}},
        "EarningsPerShareDiluted": {"units": {"USD/shares": _qs(nv_ends, [0.78, 0.89, 0.76, 1.08,
                                                                          1.30, 1.76, 2.39, 2.46])}},
        "Revenues": {"units": {"USD": _qs(nv_ends, [35e9, 39e9, 44e9, 47e9,
                                                    57e9, 68e9, 81e9, 90e9])}},
        "StockholdersEquity": {"units": {"USD": [{"end": "2026-07-26", "val": 229.0e9,
                                                  "filed": "2026-08-25"}]}}}}}
    fti = {"slim_version": 3, "facts": {"us-gaap": {
        "NetIncomeLoss": {"units": {"USD": _qs(ft_ends, [0.2e9, 0.2e9, 0.2e9, 0.25e9,
                                                         0.28e9, 0.27e9, 0.30e9, 0.33e9], lag=24)}},
        "StockholdersEquity": {"units": {"USD": [
            {"end": "2026-03-31", "val": 3.20e9, "filed": "2026-04-24"},
            {"end": "2026-06-30", "val": 1.18e9 / 0.359, "filed": "2026-07-24"}]}}}}}
    os.makedirs(os.path.join(tmp, "valuation_cache"))
    with open(os.path.join(tmp, "valuation_cache", "cik_map.json"), "w") as fh:
        json.dump({"NVDA": "0001045810", "FTI": "0001681459"}, fh)
    for cik, body in (("0001045810", nvda), ("0001681459", fti)):
        with open(os.path.join(tmp, "valuation_cache", "companyfacts_%s.json" % cik), "w") as fh:
            json.dump(body, fh)
    with open(os.path.join(tmp, "earnings_calendar_2026-10-14.json"), "w") as fh:
        json.dump([{"ticker": "NVDA", "date": "2026-11-17", "hour": "amc", "sessions": 24, "note": "",
                    "expectations": {"eps_estimate": 2.10, "beats": 4, "reported": 4, "adj_ttm": 7.0,
                                     "adj_ttm_if_met": 7.7, "chg_if_met": 0.10, "why": None}}], fh)
    txt, snap, _ = build(dt.date(2026, 10, 14), tmp, log, None, None)
    cn = {c["ticker"]: c for c in snap["card"]}
    check("end to end (DP-V4): NVDA is NEUTRAL on 12-month ROE, DEMANDING on last-FY, and BUY 74 "
          "does not flip", cn["NVDA"]["cls"] == "NEUTRAL" and cn["NVDA"]["cls_lastfy"] == "DEMANDING"
          and cn["NVDA"]["dpv2_flip"] is None and snap["card"][0]["ticker"] == "NVDA")
    check("end to end: FTI is UNDEMANDING on 12-month ROE and BUY 68 flips REJECT->ADMIT",
          cn["FTI"]["cls"] == "UNDEMANDING" and "FLIP REJECT->ADMIT" in txt)
    check("the basis line names the classes that moved off last-FY",
          "NVDA DEMANDING->NEUTRAL" in txt and "FTI NEUTRAL->UNDEMANDING" in txt)
    check("2b opens the hood: NVDA's 12-month EPS and its growth against a year earlier",
          "UNDER THE HOOD" in txt and abs(cn["NVDA"]["under_hood"]["growth"]["eps"]["ttm"] - 7.91) < 1e-9
          and cn["NVDA"]["under_hood"]["growth"]["eps"]["chg"] > 1.0)
    # ---- 2026-10-29: the cells the 10-28 sheet left blank ----
    def _u(eps_chg=0.1, ocf_ni=1.2, ocfm=0.15, loss=False, eq=100e9, le=1.5, la=0.6, cav=None,
           notes=()):
        return dict(caveat=cav, notes=list(notes),
                    growth=dict(eps=dict(ttm=2.0, chg=eps_chg), revenue=dict(chg=0.05),
                                q_eps=dict(chg=0.02)),
                    quality=dict(ocf_to_ni=ocf_ni, ocf_margin=ocfm, net_loss=loss, gross_margin=0.4,
                                 gross_margin_year_ago=0.38, op_margin=0.2, op_margin_year_ago=0.18),
                    balance=dict(equity=eq, negative_equity=eq is not None and eq <= 0,
                                 liab_to_equity=le, liab_to_assets=la, cash=1e9))
    rb = hood_sheet_row("BBWI", _u(eq=-1.056e9, le=None, la=1.2048,
                                   notes=["liabilities: derived as assets minus equity"]))
    check("sheet (BBWI): L/E blank, liabilities / assets 120%, and the note the sheet used to drop",
          rb[HOOD_SHEET_COLUMNS.index("Liabilities / equity")] is None
          and abs(rb[HOOD_SHEET_COLUMNS.index("Liabilities / assets %")] - 120.48) < 1e-9
          and rb[-1].startswith("negative book equity (deficit $1.06B): L/E not defined, see L/A")
          and rb[-1].endswith("liabilities: derived as assets minus equity"))
    rd = hood_sheet_row("DOW", _u(ocf_ni=None, ocfm=0.107, loss=True))
    check("sheet (DOW): a net loss leaves OCF/NI blank, shows OCF / revenue 10.7%, and says why",
          rd[HOOD_SHEET_COLUMNS.index("OCF / NI")] is None
          and abs(rd[HOOD_SHEET_COLUMNS.index("OCF / revenue %")] - 10.7) < 1e-9
          and "12-month net loss: OCF/NI not defined" in rd[-1])
    rs = hood_sheet_row("SND", _u(eps_chg=12.0))
    check("sheet (SND): a change beyond +/-300% prints as the card does, not as 1200",
          rs[HOOD_SHEET_COLUMNS.index("EPS y/y %")] == ">+300%" and growth_cell(-4.0) == "<-300%"
          and abs(growth_cell(0.2026) - 20.26) < 1e-9 and growth_cell(None) is None)
    rj = hood_sheet_row("JPM", _u(cav="bank: operating cash flow and liabilities/equity are not comparable"))
    check("sheet (JPM): a bank's OCF ratios and leverage stay blank, the caveat leads the note",
          all(rj[HOOD_SHEET_COLUMNS.index(k)] is None for k in
              ("OCF / NI", "OCF / revenue %", "Liabilities / equity", "Liabilities / assets %"))
          and rj[-1].startswith("bank:"))
    check("the sheet row and its header have the same width", len(rb) == len(HOOD_SHEET_COLUMNS))
    # ---- DP-V6 (2026-10-30): ROIC on 2b ----
    ri = HOOD_SHEET_COLUMNS.index("ROIC %")
    upm = dict(_u(eq=-7.94e9, le=None, la=1.2), roic=dict(
        roic=0.366, ebit=16.66e9, ebit_route="operating income", tax_basis="own 12-month effective rate",
        ic_basis="average of opening and closing", notes=[]))
    rpm = hood_sheet_row("PM", upm)
    check("ROIC (PM): 36.6% reads where ROE cannot (negative equity), with no note for a standard read",
          abs(rpm[ri] - 36.6) < 1e-9 and "ROIC:" not in rpm[-1])
    udw = dict(_u(), roic=dict(roic=0.010, ebit=0.43e9, ebit_route="pretax income + interest expense "
                               "(InterestExpenseDebt)", tax_basis="statutory fallback 25.5%",
                               ic_basis="average of opening and closing", notes=[]))
    check("ROIC (DOW): the note names the EBIT route and the 25.5% fallback",
          "ROIC: EBIT = pretax income + interest expense; tax statutory fallback 25.5%"
          in hood_sheet_row("DOW", udw)[-1])
    ujp = dict(_u(cav="bank: operating cash flow and liabilities/equity are not comparable"),
               roic=dict(roic=None, why="no 12-month operating income, and no pretax income with "
                         "interest to rebuild it"))
    check("ROIC (JPM): a bank's ROIC stays blank and adds no note beside the bank caveat",
          hood_sheet_row("JPM", ujp)[ri] is None and "ROIC:" not in hood_sheet_row("JPM", ujp)[-1])
    ugm = dict(_u(), roic=dict(roic=None, why="no current debt tag at 2026-09-30 (absence is not zero)"))
    check("ROIC (GM): no read is blank, and the note says why",
          hood_sheet_row("GM", ugm)[ri] is None and "ROIC: no current debt tag" in hood_sheet_row("GM", ugm)[-1])
    check("2b text carries the ROIC column and its legend",
          " ROIC " in txt and "ROIC (DP-V6)" in txt and "25.5%" in txt)
    check("2b text carries the OCF/rev and L/A columns and their legend",
          "OCF/rev" in txt and " L/A " in txt and "defined when equity is negative" in txt)
    check("2c: the consensus and the class if it is met (ratio / 1.10), labelled approximate",
          cn["NVDA"]["expectations"]["cls_if_met"] == "NEUTRAL"
          and abs(cn["NVDA"]["expectations"]["ratio_if_met"] - cn["NVDA"]["ratio"] / 1.10) < 1e-9
          and "APPROXIMATE" in txt)
    check("the identity check still reads implied_return's own (last-FY) ratio",
          snap["identity"]["checked"] == 2 and not snap["identity_outside_bound"])
    bs = TTMBasis(tmp)
    raw, _ = parse_exhibit(p)
    on = rebase(raw, bs, dt.date(2026, 7, 24), strict=False)["FTI"]
    before = rebase(raw, bs, dt.date(2026, 7, 24), strict=True)["FTI"]
    check("point in time: on the day FTI's 10-Q is filed the card reads it; the ledger (strictly "
          "before the session) does not", on["ttm_end"] == "2026-06-30" and before["ttm_end"] == "2026-03-31")
    fl, _ = dpv2_ledger([(dt.date(2026, 10, 15), "FTI", 68), (dt.date(2026, 10, 26), "FTI", 68)],
                        {dt.date(2026, 10, 15): {"FTI": "UNDEMANDING"},
                         dt.date(2026, 10, 26): {"FTI": "UNDEMANDING"}}, label=_label_ttm)
    check("ledger rows before DP-V4 are REBUILT, from it LIVE",
          [x[5] for x in fl] == ["rebuilt", "live"])
    tb, sb, fb = build(dt.date(2026, 10, 14), bare, log, None, None)
    cb = {c["ticker"]: c for c in sb["card"]}
    check("NEGATIVE: with no slim cache a name is NO READ and says why, never a last-FY fallback",
          cb["NVDA"]["cls"] == "NO READ" and (cb["NVDA"]["no_read_reason"] or "").startswith("no 12-month")
          and any("cik_map" in x for x in fb))
    try:
        build(dt.date(2026, 10, 15), tmp, log, None, None)
        check("NEGATIVE: a date with no exhibit refuses", False)
    except SystemExit as e:
        check("NEGATIVE: a date with no exhibit refuses", isinstance(e.code, tuple) and e.code[1] == 2)
    check("NEGATIVE: the report copy is refused as --tracker",
          main(["--date", "2026-10-14", "--root", tmp, "--log", log, "--tracker",
                os.path.join(tmp, REFUSED), "--md", os.path.join(tmp, "x.md")]) == 2)
    # ---- the embed: guards first, then the write ----
    import openpyxl
    copy = os.path.join(tmp, "tracker_with_registry.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Signals"
    wb["Signals"].cell(5, 1, dt.datetime(2026, 10, 14))
    wb["Signals"].cell(5, 2, "NVDA")
    wb["Signals"]["K5"] = "=1+1"                       # a formula must survive the embed
    wb.create_sheet("Monitor Strip")
    wb.create_sheet("How to Use")
    wb.save(copy)
    snap_j = json.loads(json.dumps(snap, default=str))
    ok, msg = embed(snap_j, copy)
    names = openpyxl.load_workbook(copy).sheetnames
    check("embed writes the sheet right after Monitor Strip",
          ok and names.index(VC_SHEET) == names.index("Monitor Strip") + 1)
    ok2, _ = embed(snap_j, copy)
    _vs = openpyxl.load_workbook(copy)[VC_SHEET]
    _hdr = [row for row in _vs.iter_rows(values_only=True) if row and row[0] == "Ticker"
            and "Liabilities / assets %" in row]
    check("the embedded sheet's 2b header carries the two new columns",
          len(_hdr) == 1 and "OCF / revenue %" in _hdr[0] and "ROIC %" in _hdr[0]
          and _hdr[0][len(HOOD_SHEET_COLUMNS) - 1] == "Notes")
    check("re-embedding replaces, never stacks",
          ok2 and openpyxl.load_workbook(copy).sheetnames.count(VC_SHEET) == 1)
    check("formulas in the copy survive the embed",
          openpyxl.load_workbook(copy)["Signals"]["K5"].value == "=1+1")
    check("no temporary file is left behind", not os.path.exists(copy + ".vc-tmp"))
    owner = os.path.join(tmp, "~$tracker_with_registry.xlsx")
    open(owner, "w").close()
    ok3, msg3 = embed(snap_j, copy)
    check("NEGATIVE: an Excel owner file refuses the write", not ok3 and "OWNER FILE" in msg3)
    os.remove(owner)
    global _probe_write
    real = _probe_write

    def _held(path):
        raise PermissionError(13, "in use")
    _probe_write = _held
    try:
        before = os.path.getmtime(copy)
        ok4, msg4 = embed(snap_j, copy)
        check("NEGATIVE: a workbook open in Excel refuses the write, file untouched",
              not ok4 and "IN USE" in msg4 and os.path.getmtime(copy) == before)
    finally:
        _probe_write = real
    stale = dict(snap_j, date="2026-10-15")
    ok5, msg5 = embed(stale, copy)
    check("NEGATIVE: a copy the review has not rebuilt today is refused", not ok5 and "latest cohort" in msg5)
    ok6, msg6 = embed(snap_j, os.path.join(tmp, "claude_equity_bot_tracker.xlsx"))
    check("NEGATIVE: the master is never written", not ok6 and "MASTER" in msg6)
    # ---- 2026-10-23: the context note --------------------------------------
    cx = {"ttm": {"PAA": {"no_read": "no TTM: no EPS tag"},
                  "DOW": {"eps_ttm": -4.0, "flags": ["STALE: last quarter ends 2026-03-31"]},
                  "SPY": {"no_read": "fund (no EPS)"}},
          "earn": {"DOW": {"date": "2026-10-22", "sessions": 0, "note": "INSIDE 5 SESSIONS"},
                   "SPY": {"note": "n/a (fund)"}}}
    check("a NO READ reason is not doubled ('TTM: no TTM:')",
          _ctx_cells(cx, "PAA")[4] == "no TTM: no EPS tag")
    check("the calendar flag and a TTM flag both print",
          _ctx_cells(cx, "DOW")[4] == "INSIDE 5 SESSIONS; TTM STALE: last quarter ends 2026-03-31")
    check("a fund says so once", _ctx_cells(cx, "SPY")[4] == "n/a (fund)")
    import rulings
    bk = rulings.RULINGS["DP-3C-bk"]
    mo = (bk.get("measurement_only") or ("ZZZZ",))[0]
    check("a DP-3C-bk name carries the thesis label while the ruling is open",
          not bk.get("open")
          or (thesis_label(mo) == THESIS_LABEL and THESIS_LABEL in _ctx_cells({}, mo)[4]))
    check("a name outside DP-3C-bk carries no label", thesis_label("AAPL") == "")
    print("SELFTEST %s (%d failed)" % ("PASSED" if not fails else "FAILED", len(fails)))
    return 1 if fails else 0


# --------------------------------------------------------------------------- #

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--date", help="YYYY-MM-DD (default: today)")
    ap.add_argument("--root", default=".")
    ap.add_argument("--tracker", help="the MASTER (logs\\claude_equity_bot_tracker.xlsx)")
    ap.add_argument("--log", help="logs\\signals_<date>.log")
    ap.add_argument("--quotes", default="quote_dump.json")
    ap.add_argument("--md", help="output exhibit (default value_card_<date>.md)")
    ap.add_argument("--json", help="also write a machine-readable snapshot here")
    ap.add_argument("--embed-from", help="a value_card_<date>.json snapshot to embed")
    ap.add_argument("--embed-into", help="the GENERATED report copy (tracker_with_registry.xlsx)")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return _selftest()
    if a.embed_from or a.embed_into:
        if not (a.embed_from and a.embed_into):
            print("COULD NOT RUN: --embed-from and --embed-into go together.")
            return 2
        if not os.path.exists(a.embed_from):
            print("COULD NOT RUN: snapshot %s not found -- step 5 did not write it." % a.embed_from)
            return 2
        with open(a.embed_from, encoding="utf-8") as fh:
            snap = json.load(fh)
        ok, msg = embed(snap, a.embed_into, a.tracker)
        print("  " + msg)
        return 0 if ok else 2
    date = dt.date.fromisoformat(a.date) if a.date else dt.date.today()
    if a.tracker and os.path.basename(a.tracker).lower() == REFUSED:
        print("REFUSED: %s is the GENERATED REPORT COPY, never an input. Pass the MASTER." % REFUSED)
        return 2
    log = a.log or os.path.join(a.root, "logs", "signals_%s.log" % date)
    if not os.path.exists(log):
        print("COULD NOT RUN: no signals log at %s" % log)
        return 2
    if a.tracker and not os.path.exists(a.tracker):
        print("COULD NOT RUN: master not found at %s" % a.tracker)
        return 2
    try:
        text, snap, findings = build(date, a.root, log, a.quotes, a.tracker)
    except SystemExit as e:
        msg, code = e.code if isinstance(e.code, tuple) else (str(e.code), 2)
        print("COULD NOT RUN: %s" % msg)
        return code
    md = a.md or os.path.join(a.root, "value_card_%s.md" % date)
    with open(md, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("# value_card exhibit\n\n```\n" + text + "\n```\n")
    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump(snap, fh, indent=1, default=str)
    print(text)
    return 1 if findings else 0


if __name__ == "__main__":
    sys.exit(main())
