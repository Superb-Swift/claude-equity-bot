#!/usr/bin/env python3
"""
outcome_sanity.py — magnitude guard for the MANUAL-ENTRY outcome columns.

WHY THIS EXISTS
---------------
2026-08-07: WMT's closing price was entered as 11.85 instead of 111.85 (dropped
leading digit). It corrupted THREE cohorts at once, because every trading day is
simultaneously the +5d, +10d and +20d resolution date for three different cohorts.

Both existing gates ran and BOTH PASSED with the bad value in place:
  * recalc.py            -> 0 errors. 11.85 is a valid number; recalc catches
                            FORMULA errors, not MAGNITUDE errors.
  * log_tracker_crosscheck.py -> 24/24 CLEAN. It reconciles columns A-G, the
                            PASTED block. Columns H/I/J are manual entry and it
                            is structurally incapable of seeing them.

The outcome columns are the only manual surface in the project with no gate,
and they carry a 3x blast radius per cell. This closes that hole.

SCOPE / DISCIPLINE
------------------
  * Reads logs\\claude_equity_bot_tracker.xlsx (THE MASTER). Writes nothing.
  * HARD-REFUSES tracker_with_registry.xlsx as an input (exit 2), same guard as
    log_tracker_crosscheck.py.
  * Formula-cache immune: reads only Price columns G/H/I/J, never K/L/M/N.
  * Uses `trading_calendar.py` for the trading-day index -- NOT the master's
    own cohort dates. AMENDED 2026-08-26: deriving the index from cohort dates
    is exact only while the bot runs every session, and it relabels itself in
    silence when one is missed (2026-08-28 and 2026-08-31 are pre-registered
    no-run sessions). The old note called the absence of a second source a
    virtue; the honest answer is to have one and RECONCILE it, which
    trading_calendar.reconcile_against_cohorts() does on every run. Still no
    network and no data file.
  * Exit 0 = clean. Exit 1 = findings. Exit 2 = wrong input / usage error.

CHECKS
------
  A  QUOTE-vs-CLOSE (primary).  A +Nd price for resolution date D is a CLOSE for
     date D. The master already holds that ticker's Px@signal for date D in the D
     cohort. Same instrument, same day: they cannot differ much. This is the check
     that catches 11.85 against 112.01.
  B  CROSS-HORIZON AGREEMENT.  Several cohorts resolve on the same date for the
     same ticker (+5d of C, +10d of C-5, +20d of C-15). Those cells must hold the
     SAME price. Catches single-cell typos. Free.
  C  MAGNITUDE BACKSTOP.  |return| >= MAG_PCT at any horizon. Needs no calendar,
     so it still fires if A and B are unavailable. A cell corroborated outside
     this check may be retired via VERIFIED_REAL_MOVES; see that block. The
     retirement is Check-C-only and is pinned to the verified price.
  D  COHORT DISPERSION.  sd(pop) of a resolved cohort > SD_PCT. A corrupted row
     blows up dispersion before it changes the mean.
  F  LOG-CLOSE RECONCILIATION (added 2026-08-26). Every session log carries
     `H1[...; closes=a>b>c>d>e]` per ticker -- a machine-sourced close series
     from Schwab. Check F compares the manual outcome cells against it. It is
     the ONLY check here with an independent price source, so it is the only
     one that can see a value propagated CONSISTENTLY into all three horizon
     cells -- the shape that produced the AG defect, where B is unanimous and
     wrong and A is inside tolerance. Reports; does not gate (see
     CHECK_F_GATES_EXIT).
     AMENDED 2026-09-16 (DP-3C-ag): F now compares EVERY filled cell against
     the log close, including dates where the cells disagree with each other
     (previously skipped as "Check B's problem"), and its tolerance is the
     log's print precision (2dp -> 0.0051), not 0.05%. The 0.05% tolerance let
     a known-wrong MSFT cell (497.12 vs 496.914, 0.04%) pass both B and F.
  G  UNSETTLED / FUTURE CELLS (added 2026-09-16, DP-3C-ag). A filled cell whose
     resolution session has not settled yet cannot hold a close. GATES. Before
     this check such cells were skipped without a word: on 2026-09-16 the
     master held 80 closes for 2026-09-17/18 and the run reported CLEAN.

Advisory-only - no orders - not financial advice.
"""

import argparse
import datetime
import os
import statistics as st
import sys

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not available. Activate the venv.", file=sys.stderr)
    sys.exit(2)

# The trading-day index. Degrades to the old cohort-derived index rather than
# failing closed -- a gate that refuses to run is a gate nobody runs -- but it
# says so loudly, because the degraded mode is the one that mis-maps horizons
# across a missed session.
try:
    import trail_basis_audit as _tba
except ImportError:                                    # pragma: no cover
    _tba = None
try:
    import trading_calendar as _cal
except ImportError:                                    # pragma: no cover
    _cal = None

# Check F reports but does not gate while GM 2026-08-20 is unadjudicated
# (manual-entry error vs corporate-action-adjusted close). Flip to True once
# that is settled; a known-open item should not turn the gate red daily.
CHECK_F_GATES_EXIT = False
CHECK_F_TOL_PCT = 0.05      # retained for the %-gap column only (see CHECK_F_TOL_ABS)
# AMENDED 2026-09-16 (DP-3C-ag). The log prints closes to 2dp, so a cell that
# holds the same close can differ from it by at most half a cent. Anything
# beyond that is a different number. The old 0.05% tolerance was 25 cents on a
# $500 stock and passed MSFT 497.12 against 496.914.
CHECK_F_TOL_ABS = 0.0051

# ---------------------------------------------------------------- thresholds
QUOTE_CLOSE_PCT = 10.0   # Check A: quote-to-close gap that warrants a look
AGREE_PCT       = 0.05   # Check B: tolerance for rounding on the same close
# Check C is horizon-scaled. Calibrated 2026-08-10 against the full record:
# the largest LEGITIMATE moves are CBRS -34.8% and MSFT +31.7% at +20d, so a
# flat 25% buried the real defect in true positives. These sit clear of both.
MAG_PCT = {"+5d": 30.0, "+10d": 40.0, "+20d": 50.0}
# The HOLD band, for the Check E breach split ONLY. This module does not and
# must not re-baseline it (Q4-2 parked); it is read here, never written.
BAND_PCT = 3.0
SD_PCT          = 12.0   # Check D: cohort dispersion ceiling

# ------------------------------------------------------- Check C allowlist
# VERIFIED REAL MOVES. Added 2026-08-26.
#
# WHY THIS EXISTS. Check C is the calendar-free backstop: it fires on magnitude
# alone so it still works when A and B cannot run. That makes it the only check
# here with no second opinion, and therefore the only one that produces standing
# true positives - a genuinely violent name reprints the same finding every day
# until the cell ages out of the +20d window. A standing finding trains the
# reader to skim the block, which defeats every other check in the file.
#
# The alternative was raising MAG_PCT. REJECTED: the thresholds were calibrated
# 2026-08-10 to sit just clear of the largest LEGITIMATE moves on record, so
# raising them re-buries the defect class the module was built to catch. An
# allowlist retires one verified cell and leaves the threshold where it is.
#
# DISCIPLINE - all four hold, or the entry does not suppress:
#   1. Scope is Check C ONLY. [A], [B] and [D] are never suppressed. If a cell
#      is allowlisted here and later disagrees across horizons or against its
#      same-day quote, it still fires - loudly, and by another name.
#   2. The entry PINS THE PRICE it was verified at. Change the cell and the
#      entry stops matching, so the finding comes back. An allowlist entry can
#      never become a permanent blind spot on a cell whose value has moved.
#   3. Ticker, row, date, horizon and price must ALL match. Row numbers shift
#      when cohort size changes (24 -> 25 -> 26 over Aug 2026); a shifted row
#      fails the match and the finding reappears. Fail-safe by construction.
#   4. Nothing is silently dropped. Suppressed findings print in their own
#      block with their reason, and the entry count prints in the header.
#
# An entry means "a human looked at this cell and corroborated it against
# something outside Check C." NOT "this number is large and I am tired of it."
#
# key   : (TICKER, master row)
# value : (cohort date, horizon, verified price, reason)
VERIFIED_REAL_MOVES = {
    ("CBRS", 1179): (
        "2026-07-29", "+10d", 262.06,
        "2026-08-26 triage. Corroborated twice outside Check C: all three "
        "cohorts resolving 2026-08-12 hold 262.06 (Check B clean), and CBRS "
        "Px@signal 2026-08-12 = 251.93, a +4.0% quote-to-close gap (Check A "
        "clean). Real move in a MEDIUM-DQ high-vol name; same ticker printed "
        "-14.60% trail5d on 2026-08-26. MAG_PCT was calibrated for CBRS "
        "-34.8% at +20d, never for +46.6% at +10d.",
    ),
}

# Price equality for the pin. Tight enough that a real re-entry breaks the
# match; loose enough to survive float round-trip through openpyxl.
PIN_TOL = 1e-6

DATE_COL, TICKER_COL, SIGNAL_COL, PXSIG_COL = 1, 2, 4, 7
HORIZONS = [("+5d", 8, 5), ("+10d", 9, 10), ("+20d", 10, 20)]
FIRST_DATA_ROW = 5

REFUSED = "tracker_with_registry.xlsx"


def load(path):
    if os.path.basename(path).lower() == REFUSED.lower():
        print("=" * 78)
        print("REFUSED: %s is the GENERATED REPORT COPY, never an input." % REFUSED)
        print("Run against logs\\claude_equity_bot_tracker.xlsx (THE MASTER).")
        print("=" * 78)
        sys.exit(2)
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Signals" not in wb.sheetnames:
        print("ERROR: no 'Signals' tab in %s" % path, file=sys.stderr)
        sys.exit(2)
    ws = wb["Signals"]
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        d = ws.cell(r, DATE_COL).value
        if not isinstance(d, datetime.datetime):
            continue
        rows.append({
            "row": r,
            "date": d.date(),
            "tkr": ws.cell(r, TICKER_COL).value,
            "sig": ws.cell(r, SIGNAL_COL).value,
            "g": ws.cell(r, PXSIG_COL).value,
            "h": ws.cell(r, 8).value,
            "i": ws.cell(r, 9).value,
            "j": ws.cell(r, 10).value,
        })
    return rows


def _digit_signature(master, log_val):
    """Is the disagreement a SINGLE MIS-KEYED DIGIT?

    ADDED 2026-08-26, after Check F's first full-archive run returned 19
    candidates and the raw list was unreadable as a queue.

    THIS IS THE TRIAGE AXIS, AND MAGNITUDE IS NOT. The two live explanations
    for a Check F hit are a manual-entry error and a corporate-action-adjusted
    close from Schwab's price history. An ADJUSTMENT is proportional and
    non-round; it cannot produce "one digit off by one". So a single-digit
    signature is strong evidence for the typo class regardless of size --
    on the first full run these ranged from 0.061% (CORN) to 2.351% (NKE),
    which is why sorting by percentage buried the clearest cases.

    Returns (True, description) or (False, "").
    """
    for dp in (2, 3, 4):
        a = ("%%.%df" % dp % master).replace(".", "")
        b = ("%%.%df" % dp % log_val).replace(".", "")
        if len(a) != len(b):
            continue
        diff = [i for i, (x, y) in enumerate(zip(a, b)) if x != y]
        if len(diff) == 1:
            i = diff[0]
            return True, "SINGLE-DIGIT %s->%s at place %d of %s" % (a[i], b[i], i + 1, a)
        # TRANSPOSITION, added 2026-09-08. Two positions differ and each holds
        # the other's digit -- the second most common keying error after a
        # single mis-key, and previously invisible to this classifier.
        #
        # THE CASE THAT FOUND IT: WMT row 1337 held 180.4200 for 2026-09-03
        # while rows 1583 and 1711 held 108.4200 for the same date. 1-0-8 keyed
        # as 1-8-0. It surfaced through Check A (a +67.5% quote-to-close gap)
        # and Check B (cross-cohort disagreement) -- both MAGNITUDE channels --
        # and through nothing that looks at the SHAPE of the error.
        #
        # That ordering is luck, not design. A transposition of adjacent
        # low-order digits produces a SMALL percentage gap: it would clear
        # Check A's 10% tolerance, and would then need the other two cohorts to
        # hold a different value to surface at all. Keyed consistently across
        # all three cells, it would be invisible to every check in this module.
        # The same argument that justifies the single-digit axis justifies this
        # one -- a corporate-action adjustment is proportional and cannot
        # transpose two digits either.
        #
        # Adjacent swaps are labelled as such because they are the common form;
        # a non-adjacent swap is still a swap and is still reported.
        if len(diff) == 2:
            i, j = diff
            if a[i] == b[j] and a[j] == b[i]:
                kind = "ADJACENT " if j == i + 1 else ""
                return True, ("%sTRANSPOSITION %s%s->%s%s at places %d-%d of %s"
                              % (kind, a[i], a[j], b[i], b[j], i + 1, j + 1, a))
    return False, ""


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _wrap(text, width):
    """Minimal greedy wrap. stdlib textwrap would do, but this module has held
    to a no-surprises import set and the reasons are short."""
    out, line = [], ""
    for word in text.split():
        if line and len(line) + 1 + len(word) > width:
            out.append(line)
            line = word
        else:
            line = word if not line else line + " " + word
    if line:
        out.append(line)
    return out


def main():
    ap = argparse.ArgumentParser(description="Magnitude guard for manual outcome columns.")
    ap.add_argument("--tracker", default=os.path.join("logs", "claude_equity_bot_tracker.xlsx"))
    ap.add_argument("--quiet", action="store_true", help="findings only")
    ap.add_argument("--logs", default=os.path.join("logs", "signals_*.log"),
                    help="session logs for Check F's close reconciliation")
    ap.add_argument("--quote-close-pct", type=float, default=QUOTE_CLOSE_PCT)
    ap.add_argument("--mag-scale", type=float, default=1.0,
                    help="multiplier on the horizon-scaled magnitude thresholds")
    ap.add_argument("--sd-pct", type=float, default=SD_PCT)
    ap.add_argument("--shape-n", type=int, default=10,
                    help="how many recent cohorts to show in the Check E shape table")
    ap.add_argument("--no-allowlist", action="store_true",
                    help="ignore VERIFIED_REAL_MOVES and show the raw ungated "
                         "[C] set (audit view - run this before trusting a clean result)")
    args = ap.parse_args()

    if not os.path.exists(args.tracker):
        print("ERROR: not found: %s" % args.tracker, file=sys.stderr)
        sys.exit(2)

    findings_cal = []
    rows = load(args.tracker)
    dates = sorted({r["date"] for r in rows})
    idx = {d: i for i, d in enumerate(dates)}          # legacy fallback only

    # AMENDED 2026-08-26. `resolve(d, n)` is the +n-trading-day resolution date
    # and is correct whether or not the bot ran on any session in between.
    cal_notes = []
    if _cal is not None:
        def resolve(d, n):
            return _cal.shift_sessions(d, n)
        rec = _cal.reconcile_against_cohorts(dates)
        for d in rec["missing"]:
            findings_cal.append("[CAL] %s is an NYSE session with no cohort row "
                                "and no BOT_ABSENT entry. Either the bot missed a "
                                "run or the calendar is wrong; horizons around it "
                                "cannot be trusted until it is explained." % d)
        for d in rec["extra"]:
            findings_cal.append("[CAL] cohort rows exist for %s, which the calendar "
                                "says is not an NYSE session." % d)
        for d in rec["absent_expected"]:
            why, doc = _cal.BOT_ABSENT[d]
            cal_notes.append("%s - %s (%s)" % (d, why, doc))
    else:                                              # pragma: no cover
        def resolve(d, n):
            i = idx[d]
            return dates[i + n] if i + n < len(dates) else None
        cal_notes.append("trading_calendar.py NOT IMPORTABLE - horizons are "
                         "derived from cohort dates and WILL mis-map across any "
                         "missed session.")
    last_cohort = dates[-1]
    # (date, ticker) -> Px@signal, our same-day reference price
    pxsig = {(r["date"], r["tkr"]): r["g"] for r in rows if num(r["g"])}

    findings = []
    checked_a = skipped_a = 0
    after_cohort = 0          # filled cells resolving after the newest cohort

    # ---- Check G: a filled cell must resolve on a SETTLED session -----------
    # ADDED 2026-09-16 (DP-3C-ag). See the header. `settled_through` is the
    # last session whose regular close is final on this machine's clock; the
    # 16:15 ET rule is repair_outcome_close.market_settled(), one rule for the
    # project. If that module is unavailable the newest cohort date is used,
    # which is stricter (it refuses an evening entry for a session the bot did
    # not run) and never looser.
    findings_g = []
    settled_through = last_cohort
    settled_src = "newest cohort date (settlement clock unavailable)"
    try:
        from repair_outcome_close import market_settled as _ms
        if _cal is not None:
            t = datetime.date.today()
            if _cal.is_trading_day(t) and _ms(t)[0]:
                settled_through = t
            else:
                settled_through = _cal.prior_session(t)
            settled_src = "16:15 ET settlement clock"
    except Exception:                                   # pragma: no cover
        pass
    for r in rows:
        for label, col, n in HORIZONS:
            v = r[{8: "h", 9: "i", 10: "j"}[col]]
            if not num(v):
                continue
            rdate = resolve(r["date"], n)
            if rdate is None:
                findings_g.append("[G] %-5s row %-5d %s %s holds %.4f but its "
                                  "resolution date cannot be computed"
                                  % (r["tkr"], r["row"], r["date"], label, v))
            elif rdate > settled_through:
                findings_g.append("[G] %-5s row %-5d %s %s holds %.4f for %s, "
                                  "which has NOT SETTLED (settled through %s)"
                                  % (r["tkr"], r["row"], r["date"], label, v,
                                     rdate, settled_through))

    # ---- Check A: quote-vs-close, and Check B: cross-horizon agreement ------
    resolved_at = {}   # (resolution_date, ticker) -> list of (label, row, price)
    for r in rows:
        for label, col, n in HORIZONS:
            v = r[{8: "h", 9: "i", 10: "j"}[col]]
            if not num(v):
                continue
            rdate = resolve(r["date"], n)
            if rdate is None or rdate > last_cohort:
                # No same-day Px@signal exists yet, so A cannot run. This used
                # to be a SILENT skip; the count is now printed, and a cell
                # whose session has not settled is a Check G finding above.
                after_cohort += 1
                continue
            resolved_at.setdefault((rdate, r["tkr"]), []).append((label, r["row"], v, r["date"]))

            ref = pxsig.get((rdate, r["tkr"]))
            if not num(ref) or ref == 0:
                skipped_a += 1
                continue
            checked_a += 1
            gap = (v - ref) / ref * 100.0
            if abs(gap) > args.quote_close_pct:
                findings.append(
                    "[A] %-5s row %-5d %s %s price %.4f vs Px@signal %.4f on %s "
                    "-> %+.1f%% quote-to-close gap"
                    % (r["tkr"], r["row"], r["date"], label, v, ref, rdate, gap))

    info_b = []      # sub-tolerance disagreements: reported, never gating
    for (rdate, tkr), entries in sorted(resolved_at.items()):
        vals = [e[2] for e in entries]
        if not vals:
            continue
        detail = "; ".join("%s from %s (row %d) = %.4f" % (e[0], e[3], e[1], e[2]) for e in entries)
        spread_pct = (max(vals) - min(vals)) / max(abs(min(vals)), 1e-9) * 100.0
        if spread_pct > AGREE_PCT:
            findings.append("[B] %-5s disagreeing closes for %s -> %s" % (tkr, rdate, detail))
        elif max(vals) - min(vals) > 1e-6:
            # ADDED 2026-09-16 (DP-3C-ag). Two cells for the same close that
            # differ at all are not both right. Below AGREE_PCT this is not a
            # finding (the tolerance is unchanged) but it is no longer silent:
            # MSFT 2026-09-15 (497.12 vs 496.914, 0.04%) sat here unreported.
            info_b.append("[b] %-5s %s differs by %.3f%% -> %s"
                          % (tkr, rdate, spread_pct, detail))

    # ---- Check C: magnitude backstop ---------------------------------------
    suppressed = []          # allowlisted [C] hits, reported but not gating
    matched_keys = set()     # allowlist entries that actually fired
    for r in rows:
        if not num(r["g"]) or r["g"] == 0:
            continue
        for label, col, _ in HORIZONS:
            v = r[{8: "h", 9: "i", 10: "j"}[col]]
            if not num(v):
                continue
            ret = (v - r["g"]) / r["g"] * 100.0
            if abs(ret) < MAG_PCT[label] * args.mag_scale:
                continue
            line = ("[C] %-5s row %-5d %s %s return %+.1f%% (%.4f -> %.4f)"
                    % (r["tkr"], r["row"], r["date"], label, ret, r["g"], v))

            # Retire only on a FULL match: ticker, row, date, horizon, price.
            # Any drift in any field and the finding stands. See the
            # VERIFIED_REAL_MOVES block for why each field is in the key.
            entry = None if args.no_allowlist else VERIFIED_REAL_MOVES.get((r["tkr"], r["row"]))
            if (entry
                    and entry[0] == str(r["date"])
                    and entry[1] == label
                    and abs(entry[2] - v) <= PIN_TOL):
                matched_keys.add((r["tkr"], r["row"]))
                suppressed.append((line, entry[3]))
            else:
                findings.append(line)

    # ---- Check F: reconcile outcome cells against the log's close channel ---
    # WHY THIS IS THE ONLY CHECK THAT CAN SEE THE AG DEFECT CLASS. A, B and D
    # are all internal to the master: B compares the master against itself, so a
    # value propagated consistently into all three horizon cells is unanimous
    # and passes; A compares a close against a same-day 09:3x quote, a window
    # wide enough to swallow a plausible-looking wrong number. F brings in a
    # price the operator did not type.
    #
    # CORROBORATION IS THE POINT. The same session's close appears in up to
    # four different logs, because each run carries a multi-day close series.
    # A disagreement confirmed by four independent Schwab fetches is a different
    # claim from one seen once, and the two are NOT reported as if they were.
    findings_f = []
    findings_f_prec = []     # below 0.05% but beyond the log's print precision
    checked_f = 0
    if _tba is not None:
        try:
            closes_ref = {}
            for r in rows:
                for label, col, n in HORIZONS:
                    v = r[{8: "h", 9: "i", 10: "j"}[col]]
                    if num(v):
                        rd = resolve(r["date"], n)
                        if rd is not None:
                            closes_ref.setdefault((r["tkr"], rd), []).append((v, r["row"], label))
            import glob as _glob
            logs = sorted(set(_glob.glob(args.logs) +
                              _glob.glob(os.path.join(os.path.dirname(args.logs) or ".",
                                                      "**", os.path.basename(args.logs)),
                                         recursive=True)))
            seen = {}       # (tkr, session) -> {log price -> [run dates]}
            for lp in logs:
                m = _tba.LOG_DATE_RE.search(os.path.basename(lp))
                if not m:
                    continue
                run_date = datetime.date.fromisoformat(m.group(1))
                per = _tba.parse_log(lp)
                if not per:
                    continue
                # Alignment is MEASURED per log, never assumed: if the H1 window
                # off-by-one is ever fixed, this follows it rather than breaking.
                ref_for_align = {k: v[0][0] for k, v in closes_ref.items()}
                a_in = _tba.score_alignment(per, run_date, ref_for_align, True)
                a_do = _tba.score_alignment(per, run_date, ref_for_align, False)
                last_is_run = a_in[0] >= a_do[0]
                for tkr, (_lbl, _tr, cl) in per.items():
                    k = len(cl)
                    for pos, price in enumerate(cl):
                        back = (k - 1 - pos) if last_is_run else (k - pos)
                        if back <= 0:
                            continue        # the live intraday element: no settled ref
                        sess = resolve(run_date, -back)
                        if sess is None:
                            continue
                        seen.setdefault((tkr, sess), {}).setdefault(price, []).append(run_date)

            for (tkr, sess), byprice in sorted(seen.items()):
                cells = closes_ref.get((tkr, sess))
                if not cells:
                    continue
                # The log price corroborated by the most runs is the challenger.
                price, runs = max(byprice.items(), key=lambda kv: len(kv[1]))
                spread = ("" if len(byprice) == 1 else
                          "  [logs disagree: also %s]"
                          % ", ".join("%.4f x%d" % (p, len(d))
                                      for p, d in sorted(byprice.items()) if p != price))
                # AMENDED 2026-09-16 (DP-3C-ag): EVERY cell is compared. Dates
                # whose cells disagree with each other used to be skipped as
                # "Check B's problem", so a wrong cell with a slightly different
                # sibling passed both checks. Now each cell is judged against
                # the log, and when the cells disagree the line says which one
                # the log sides with -- the adjudication hint.
                bad, good = [], []
                for v, row, label in cells:
                    checked_f += 1
                    (bad if abs(v - price) > CHECK_F_TOL_ABS else good).append((v, row, label))
                if not bad:
                    continue
                # One line per (ticker, session). Cells holding the same wrong
                # value are listed together; cells holding different wrong
                # values each get a line.
                by_val = {}
                for v, row, label in bad:
                    by_val.setdefault(round(v, 6), []).append((row, label))
                for v, where_l in sorted(by_val.items()):
                    gap = (price - v) / v * 100.0 if v else float("nan")
                    material = abs(gap) >= CHECK_F_TOL_PCT
                    is_digit, sig = _digit_signature(v, price)
                    if not material:
                        # A last-place difference always "looks" like a
                        # single mis-keyed digit; the signature carries no
                        # information at this size and is not printed.
                        is_digit, sig = False, "precision tier"
                    hint = ""
                    if good:
                        hint = "  ADJUDICATION: log agrees with %s" % "/".join(
                            "row %d %s (%.4f)" % (g[1], g[2], g[0]) for g in good)
                    where = "/".join("row %d %s" % w for w in where_l)
                    findings_f.append((
                        0 if material else 1, 0 if is_digit else 1, -len(runs), -abs(gap),
                        "[F] %-5s %s  master %.4f vs log %.4f  %+.3f%%  %d log(s)  "
                        "cells %s  %s%s%s"
                        % (tkr, sess, v, price, gap, len(runs), where,
                           sig if is_digit else ("no digit signature" if material else sig),
                           spread, hint)))
            findings_f.sort()
            findings_f_prec = [f[-1] for f in findings_f if f[0] == 1]
            findings_f = [f[-1] for f in findings_f if f[0] == 0]
        except Exception as e:                                 # pragma: no cover
            findings_f.append("[F] reconciliation unavailable: %s" % e)
    else:                                                      # pragma: no cover
        findings_f.append("[F] trail_basis_audit.py not importable - the log close "
                          "channel was not read this run.")

    # ---- Check D: cohort dispersion ----------------------------------------
    for d in dates:
        rr = []
        for r in rows:
            if r["date"] != d or not num(r["g"]) or r["g"] == 0 or not num(r["h"]):
                continue
            rr.append((r["h"] - r["g"]) / r["g"])
        if len(rr) >= 5:
            sd = st.pstdev(rr) * 100.0
            if sd > args.sd_pct:
                findings.append("[D] cohort %s +5d dispersion sd(pop) %.2f%% exceeds %.1f%%"
                                % (d, sd, args.sd_pct))

    # ---- Check E: COHORT SHAPE — off-centering and breach asymmetry ---------
    # ADDED 2026-08-13. NOT A FINDING and deliberately not appended to
    # `findings`: this reports, it does not gate.
    #
    # WHY IT BELONGS HERE. +5d "accuracy" on a ~92% HOLD population is one minus
    # the band-breach rate, and the band is pinned at 0 +/- 3% while the return
    # distribution sits at its own mean. Two numbers therefore explain almost all
    # of it, and neither is printed anywhere today:
    #   OFF-CENTERING = cohort mean / cohort sd. Measured over 56 resolved
    #     cohorts, %UP-breach against off-centering gives pearson r = +0.926.
    #     That correlation is MECHANICAL, not a discovery -- it is the band-breach
    #     identity restated -- which is exactly why it is a useful diagnostic and
    #     must never be quoted as a finding. All-56 median is 0.03; 2026-08-03
    #     through 08-06 ran 0.53-0.65, roughly twenty times the central value.
    #   UP/DOWN SPLIT of the breaches. Pooled across 56 cohorts it is 311/279 =
    #     52.7% up, near a coin flip. Over 2026-08-03..08-06 it was 30 UP / 1 DOWN.
    #     "It rose while we held" and "it fell while we held" score identically
    #     under a symmetric band; for a system whose terminal intent is execution
    #     they are not the same error.
    #
    # The E4 clock started 2026-08-13 into an unusually calm, strongly up-drifting
    # tape. If drift mean-reverts inside the E4 window, +5d accuracy falls for
    # reasons that have nothing to do with EDGAR. This block is what makes that
    # separable at the read.
    #
    # HONEST SCOPE NOTE (2026-08-13): an earlier draft of this rationale claimed
    # the figures were "unreconstructable for cohorts already past." THAT WAS
    # WRONG. Both derive from the Price columns, which are retained, so any past
    # cohort can be recomputed at any time. Printing them is a convenience and a
    # habit-former, NOT a data-preservation measure, and nothing here needed to
    # be rushed on that basis. Recorded rather than quietly dropped.
    shape = []
    for d in dates:
        vals, hold_vals = [], []
        for r in rows:
            if r["date"] != d or not num(r["g"]) or r["g"] == 0 or not num(r["h"]):
                continue
            ret = (r["h"] - r["g"]) / r["g"] * 100.0
            vals.append(ret)
            # The +/-3% band scores HOLD rows ONLY. A directional row is scored by
            # the sign test and can never "breach" - pooling it here would inflate
            # the count with rows the band was never applied to.
            if str(r["sig"]).strip().upper() == "HOLD":
                hold_vals.append(ret)
        if len(vals) < 5:
            continue
        mean, sd = st.mean(vals), st.pstdev(vals)
        breaches = [v for v in hold_vals if abs(v) > BAND_PCT]
        up = sum(1 for v in breaches if v > 0)
        shape.append((d, len(vals), mean, sd, (mean / sd if sd else float("nan")),
                      len(breaches), up, len(breaches) - up))

    if not args.quiet and shape:
        print("-" * 78)
        print("  COHORT SHAPE - context only, NOT a gate. Last %d resolved cohort(s):"
              % min(args.shape_n, len(shape)))
        print("    %-12s %4s %8s %7s %9s   %s" %
              ("cohort", "n", "mean%", "sd%", "off-ctr", "breach (UP/DOWN)"))
        for d, n_, mean, sd, oc, nb, up, dn in shape[-args.shape_n:]:
            print("    %-12s %4d %+8.2f %7.2f %9.2f   %3d (%d UP / %d DOWN)"
                  % (d, n_, mean, sd, oc, nb, up, dn))
        print("    off-ctr = mean/sd against a band fixed at 0 +/- %.0f%%. It is a"
              % BAND_PCT)
        print("    DEFINITIONAL restatement of the breach rate, not evidence.")

    # ---- COHORT COMPLETENESS -- which NAME went missing, and when -----------
    # ADDED 2026-09-08. The shape table above has been printing the evidence
    # for this since 2026-08-13 and nothing read it: cohort 2026-08-18 shows
    # n = 25 while every neighbouring cohort shows 26. The missing name is GM,
    # confirmed independently by h1_lag_trace, whose GM trace jumps 2026-08-17
    # -> 2026-08-19 while WMT carries an 08-18 row.
    #
    # GM IS ONE OF THE TWO PRE-REGISTERED WS2 TRACERS. Its episode clock reads
    # "0 post-deploy episode(s) in 42 session(s)" against WMT's 43 -- the gap
    # silently absorbed into a denominator. If the missing row carried a
    # trailing move, an episode could have been missed outright.
    #
    # WHY NOTHING CAUGHT IT. The calendar reconciliation checks SESSIONS ("no
    # unexplained missing sessions vs master", "no cohort rows on non-sessions")
    # and signal_scan checks membership against THE PREVIOUS LOG at run time.
    # Neither asks whether a cohort already in the master is short a name. A
    # per-ticker gap inside an otherwise-present cohort is invisible to every
    # gate in the chain.
    #
    # This reports and does not gate: a name legitimately leaves the watchlist,
    # and turning that into a failure would be wrong. What it must not do is
    # stay silent, which is what it did for three weeks.
    if not args.quiet and len(dates) >= 2:
        by_date = {}
        for r in rows:
            # FIXED 2026-09-16: this read r.get("ticker"), a key load() never
            # sets (it is "tkr"), so every cohort was the set {None} and the
            # detector added on 2026-09-08 could never fire. GM 2026-08-18 was
            # never reported by it.
            by_date.setdefault(r["date"], set()).add(r.get("tkr"))
        gaps = []
        ordered = sorted(by_date)
        for i in range(1, len(ordered) - 1):
            prev_s, cur_s, next_s = (by_date[ordered[i - 1]], by_date[ordered[i]],
                                     by_date[ordered[i + 1]])
            # Present either side, absent in the middle: a GAP, not a departure.
            missing = sorted((prev_s & next_s) - cur_s)
            if missing:
                gaps.append((ordered[i], missing, len(cur_s), len(prev_s)))
        if gaps:
            print("-" * 78)
            print("  COHORT COMPLETENESS - %d cohort(s) missing a name that is present"
                  % len(gaps))
            print("  on BOTH neighbours. Reported, NOT a gate: a name may legitimately")
            print("  leave the watchlist, but a name that returns the next session did")
            print("  not leave -- it was dropped, and any per-ticker denominator built")
            print("  over that window is one short without saying so.")
            for d, missing, n_cur, n_prev in gaps:
                print("    %-12s n=%d (neighbours %d)  MISSING: %s"
                      % (d, n_cur, n_prev, ", ".join(missing)))

    # ---- report -------------------------------------------------------------
    if not args.quiet:
        print("=" * 78)
        print("OUTCOME-COLUMN SANITY CHECK - %s" % datetime.date.today())
        print("=" * 78)
        print("  master  : %s" % args.tracker)
        print("  rows    : %d | cohorts %d | %s -> %s" % (len(rows), len(dates), dates[0], dates[-1]))
        print("  thresholds: quote-close %.1f%% | magnitude %s (x%.2f) | dispersion %.1f%%"
              % (args.quote_close_pct,
                 "/".join("%s %.0f%%" % (k, v) for k, v in MAG_PCT.items()),
                 args.mag_scale, args.sd_pct))
        print("  [A] quote-vs-close comparisons made: %d (skipped %d - no same-day reference)"
              % (checked_a, skipped_a))
        print("  [A/B] %d filled cell(s) resolve after the newest cohort (%s) - no"
              % (after_cohort, last_cohort))
        print("        same-day reference yet; each is checked by [G] below")
        print("  [G] settled through %s (%s)" % (settled_through, settled_src))
        print("  [C] allowlist: %d verified-real-move entry(ies)%s"
              % (len(VERIFIED_REAL_MOVES),
                 " - DISABLED (--no-allowlist)" if args.no_allowlist else ""))
        print("-" * 78)

    # ---- allowlist hygiene --------------------------------------------------
    # An entry that no longer matches anything is reported, never assumed
    # harmless: it means the cell, the row or the date moved under it. Not a
    # finding (it suppresses nothing, so it cannot hide a defect), but it is
    # dead weight in a gate and dead weight is how gates rot.
    stale = [k for k in VERIFIED_REAL_MOVES if k not in matched_keys]
    if stale and not args.no_allowlist and not args.quiet:
        print("  [!] %d STALE allowlist entry(ies) - matched no [C] hit this run:" % len(stale))
        for tkr, row in sorted(stale, key=lambda k: k[1]):
            d, lbl, px, _ = VERIFIED_REAL_MOVES[(tkr, row)]
            print("      %-5s row %-5d %s %s @ %.4f - row/date/price moved, or the"
                  % (tkr, row, d, lbl, px))
            print("      %s move aged out. Verify, then delete the entry." % (" " * 5))
        print("-" * 78)

    if suppressed:
        print("  %d SUPPRESSED [C] finding(s) - verified real moves, NOT gating:" % len(suppressed))
        for line, reason in suppressed:
            print("    " + line)
            for chunk in _wrap(reason, 68):
                print("        " + chunk)
        print("    Re-run with --no-allowlist to see these as findings.")
        print("-" * 78)

    if cal_notes and not args.quiet:
        print("  CALENDAR - pre-registered bot-absent session(s) in range:")
        for nline in cal_notes:
            print("      " + nline)
        print("      Horizons are mapped on the NYSE calendar, so these do NOT")
        print("      shift any resolution date. Recorded so the gap stays visible.")
        print("-" * 78)

    if info_b and not args.quiet:
        print("  %d SUB-TOLERANCE [b] disagreement(s) - two cells for one close that"
              % len(info_b))
        print("  differ by less than %.2f%%. INFO, not a finding; one of them is"
              % AGREE_PCT)
        print("  still wrong. Check F below says which, where a log carries it.")
        for line in info_b:
            print("    " + line)
        print("-" * 78)

    if not args.quiet:
        print("  [F] log-close comparisons made: %d cell(s), tolerance %.4f "
              "(the log's 2dp print)" % (checked_f, CHECK_F_TOL_ABS))
    if findings_f:
        print("  %d CHECK F FINDING(S) - independent price source disagrees "
              "with a manual cell%s:"
              % (len(findings_f), "" if CHECK_F_GATES_EXIT else " (REPORTING ONLY)"))
        n_digit = sum(1 for f in findings_f if "SINGLE-DIGIT" in f)
        n_trans = sum(1 for f in findings_f if "TRANSPOSITION" in f)
        if n_digit:
            print("    %d with a SINGLE MIS-KEYED DIGIT - a corporate-action adjustment"
                  % n_digit)
            print("    is proportional and cannot produce this, so these are the")
            print("    typo-class queue. Size does NOT rank them; the signature does.")
        if n_trans:
            print("    %d with a TRANSPOSITION (two digits swapped) - same argument:"
                  % n_trans)
            print("    an adjustment cannot swap two digits either. ADDED 2026-09-08")
            print("    after WMT 180.4200/108.4200 surfaced through magnitude checks")
            print("    alone; a low-order swap would have cleared all of them.")
        for f in findings_f:
            print("    " + f)
        print("    Resolve with the settled close (repair_outcome_close.py --price)")
        print("    before treating either value as correct. A corporate-action")
        print("    adjustment in Schwab's price history is a live alternative to a")
        print("    manual-entry error and would mean the two channels are not")
        print("    interchangeable.")
        print("-" * 78)
    if findings_f_prec and not args.quiet:
        print("  %d CHECK F PRECISION-TIER item(s) - the cell differs from the log's"
              % len(findings_f_prec))
        print("  2dp close by more than half a cent but by less than 0.05%. Not")
        print("  a typo queue: most are a different price source (a quote or a")
        print("  consolidated last trade) typed as the close. ADDED 2026-09-16")
        print("  (DP-3C-ag) because MSFT 2026-09-15 sat in exactly this band.")
        for f in findings_f_prec:
            print("    " + f)
        print("-" * 78)

    if findings_g:
        print("  %d CHECK G FINDING(S) - a close for a session that has not settled:"
              % len(findings_g))
        for f in findings_g[:40]:
            print("    " + f)
        if len(findings_g) > 40:
            print("    ... and %d more" % (len(findings_g) - 40))
        print("    Clear these cells. A close cannot exist before its session ends;")
        print("    enter closes with enter_closes.py, which refuses unsettled sessions.")
        print("-" * 78)

    g_summary = (["[G] %d cell(s) hold a close for an unsettled session - listed above"
                  % len(findings_g)] if findings_g else [])
    findings = findings_cal + g_summary + findings + (findings_f if CHECK_F_GATES_EXIT else [])

    if findings:
        print("  %d FINDING(S) - VERIFY BEFORE READING ANY RESOLUTION:" % len(findings))
        for f in findings:
            print("    " + f)
        print("-" * 78)
        print("  A manual outcome price is entered once and read by THREE cohorts")
        print("  (+5d, +10d, +20d all land on the same date). Fix at source in the")
        print("  master, then re-run. Do not read cohort accuracy until this is clean.")
        print("=" * 78)
        return 1

    if not args.quiet:
        if suppressed:
            print("  RESULT: CLEAN - no implausible outcome prices "
                  "(%d verified real move(s) suppressed above)." % len(suppressed))
        else:
            print("  RESULT: CLEAN - no implausible outcome prices.")
        print("=" * 78)
    return 0


if __name__ == "__main__":
    sys.exit(main())
